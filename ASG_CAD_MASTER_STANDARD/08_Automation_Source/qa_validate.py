"""
ASG CAD Master Standard - QA validation + documentation.

Re-opens every generated file from disk and tests it. Writes:
  06_Layer_Register/ASG_Layer_Register.xlsx (+ .csv)
  05_CAD_Standards_Documentation/ASG_CAD_Standards.pdf
  07_QA_Validation/ASG_Template_Validation_Report.pdf (+ .txt)
  07_QA_Validation/QA_render_*.png / .pdf

Statuses: PASS (tested successfully) | PARTIAL (part tested) |
          UNVERIFIED (could not be tested) | FAIL (test failed)
"""

import csv
import datetime as dt
import hashlib
import subprocess
import tempfile
from pathlib import Path

import ezdxf
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from ezdxf import recover  # noqa: E402
from ezdxf.addons import acadctb  # noqa: E402
from ezdxf.addons.drawing import Frontend, RenderContext  # noqa: E402
from ezdxf.addons.drawing.config import (BackgroundPolicy, ColorPolicy,  # noqa: E402
                                         Configuration, LineweightPolicy)
from ezdxf.addons.drawing.matplotlib import MatplotlibBackend  # noqa: E402

import asg_standard as S  # noqa: E402
import build_all as BA  # noqa: E402

ROOT = BA.ROOT
D_DOC = ROOT / "05_CAD_Standards_Documentation"
D_REG = ROOT / "06_Layer_Register"
D_QA = BA.D_QA
TODAY = dt.date.today().strftime("%d.%m.%Y")

R = []  # (area, test, status, evidence)


def rec(area, test, status, evidence=""):
    R.append((area, test, status, evidence))


def pf(ok):
    return "PASS" if ok else "FAIL"


def anno(e):
    try:
        return e.has_xdata("AcadAnnotative")
    except Exception:
        return False


def scale_names(doc):
    out = []
    for _k, e in doc.rootdict.get("ACAD_SCALELIST").items():
        t = e.xtags.get_subclass("AcDbScale")
        out.append((t.get_first_value(300), t.get_first_value(141)))
    return out


def render(doc, layout, png, dpi, pdf=None):
    w, h = layout.dxf.paper_width, layout.dxf.paper_height
    fig = plt.figure(figsize=(w / 25.4, h / 25.4))
    ax = fig.add_axes([0, 0, 1, 1])
    cfg = Configuration(color_policy=ColorPolicy.BLACK, lineweight_policy=LineweightPolicy.ABSOLUTE,
                        background_policy=BackgroundPolicy.WHITE)
    ctx = RenderContext(doc)
    ctx.set_current_layout(layout)
    Frontend(ctx, MatplotlibBackend(ax, adjust_figure=False), config=cfg).draw_layout(layout, finalize=True)
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)
    ax.set_aspect("auto")
    fig.savefig(png, dpi=dpi)
    if pdf:
        fig.savefig(pdf)
    plt.close(fig)


# ===========================================================================
def test_environment():
    found = [t for t in ("acad", "accoreconsole", "ODAFileConverter", "dwg2dxf", "dxf2dwg",
                         "qcad", "librecad", "bricscad") if subprocess.run(
        ["which", t], capture_output=True).returncode == 0]
    rec("Environment", "Native AutoCAD / DWG writer available", "FAIL" if not found else "PASS",
        "None of AutoCAD, accoreconsole, ODA File Converter, LibreDWG, QCAD, BricsCAD present"
        if not found else ", ".join(found))
    for ext in ("dwt", "dwg"):
        rec("Deliverables", f"ASG_Master_Template.{ext} generated natively", "FAIL",
            "Not generated - no DWG writer. Not faked by renaming. Create in AutoCAD per README step 4.")


def test_master(path):
    doc, aud = recover.readfile(path)
    rec("File", "Master DXF re-opened (recover + audit)", pf(not aud.has_errors),
        f"{len(aud.errors)} errors, {len(aud.fixes)} fixes, {doc.dxfversion} = AutoCAD 2018 DXF")
    ezdxf.readfile(path)
    rec("File", "Master DXF strict parse", "PASS", "ezdxf.readfile without exception")
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / path.name
        tmp.write_bytes(path.read_bytes())
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(tmp)],
                       capture_output=True, timeout=180)
        ok = (Path(td) / (path.stem + ".pdf")).exists()
    rec("File", "Master DXF opened by second, independent reader (LibreOffice Draw)", pf(ok),
        "converted without load error" if ok else "load failed")
    rec("File", "Opens in AutoCAD", "UNVERIFIED", "AutoCAD not available - README step 2")

    h = doc.header
    for var, val, purpose in S.HEADER_VARS + S.HEADER_STYLE_VARS:
        got = h.get(var)
        ok = (tuple(got) == tuple(val)) if isinstance(val, tuple) else got == val
        rec("Units & variables", f"{var} = {val}", pf(ok), f"read back {got} - {purpose}")
    rec("Units & variables", "Machine/profile variables (OSMODE, AUTOSNAP, PICKBOX ...)", "UNVERIFIED",
        "Not drawing-resident; applied only by opt-in ASG-PROFILE in AutoCAD")

    for name in list(S.LINETYPES) + ["Continuous"]:
        rec("Linetypes", f"{name} loaded", pf(name in doc.linetypes))
    names = [l.dxf.name for l in doc.layers]
    asg = [n for n in names if n.startswith("ASG-")]
    rec("Layers", f"{len(S.LAYERS)} ASG layers present", pf(len(asg) == len(S.LAYERS)), f"found {len(asg)}")
    rec("Layers", "No other layers except system 0 / Defpoints",
        pf(sorted(set(names) - set(asg)) == ["0", "Defpoints"]), str(sorted(set(names) - set(asg))))
    rec("Layers", "Layer names and codes unique",
        pf(len({l[1] for l in S.LAYERS}) == len(S.LAYERS) == len({l[0] for l in S.LAYERS})))
    bad = []
    for code, n, _g, aci, lt, lw, plot, desc, _u in S.LAYERS:
        l = doc.layers.get(n)
        if not (l.dxf.color == aci and l.dxf.linetype == lt and l.dxf.lineweight == round(lw * 100)
                and bool(l.dxf.plot) == plot and l.description == desc and lt in doc.linetypes):
            bad.append(n)
    rec("Layers", "Every layer: ACI, linetype (loaded), lineweight, plot flag, description",
        pf(not bad), f"{len(S.LAYERS) - len(bad)}/{len(S.LAYERS)} correct {bad if bad else ''}")
    noplot = sorted(l[1] for l in S.LAYERS if not l[6])
    rec("Layers", "No-plot layers exactly: " + ", ".join(noplot),
        pf(sorted(l.dxf.name for l in doc.layers if l.dxf.name.startswith("ASG") and not l.dxf.plot) == noplot))
    rec("Layers", "Current layer is plottable geometry",
        pf(doc.layers.get(h["$CLAYER"]).dxf.plot == 1), h["$CLAYER"])

    for n, ttf, _b, a, _u in S.TEXT_STYLES:
        st = doc.styles.get(n)
        ok = st.dxf.font == ttf and st.dxf.height == 0 and st.dxf.width == 1 and anno(st) == a
        rec("Text styles", f"{n}: {ttf}, height 0, annotative={a}", pf(ok))
    rec("Text styles", "Annotative text behaviour at viewport scales", "UNVERIFIED",
        "Annotative flag written & re-read; scaling behaviour needs AutoCAD (README step 5)")

    for n, a, over, _u in S.DIM_STYLES:
        d = doc.dimstyles.get(n)
        exp = dict(S.DIM_BASE, **over)
        diff = [k for k, v in exp.items() if k != "dimblk" and d.dxf.get(k) != v]
        rec("Dimension styles", f"{n}", pf(d is not None and not diff and anno(d) == a),
            f"all {len(exp)} variables match; annotative={a}" if not diff else f"mismatch {diff}")
    for n, a, arrow, asz, ts, ch, frame, _u in S.MLEADER_STYLES:
        m = doc.mleader_styles.get(n)
        ok = (m.dxf.char_height == ch and m.dxf.arrow_head_size == asz and bool(m.dxf.is_annotative) == a
              and m.dxf.leader_type == 1 and m.dxf.max_leader_segments_points == 2
              and bool(m.dxf.has_text_frame) == frame
              and m.dxf.text_style_handle == doc.styles.get(ts).dxf.handle)
        rec("Multileader styles", n, pf(ok), f"{ts} {ch} mm, arrow {arrow or 'closed filled'} {asz}")
    rec("Multileader styles", "Background mask", "UNVERIFIED",
        "Mask is an MText property, not stored in MLEADERSTYLE - documented per-note procedure")
    for n in ("ASG-MLINE-2", "ASG-MLINE-3"):
        rec("Multiline styles", n, pf(doc.mline_styles.get(n) is not None))

    sc = scale_names(doc)
    want = [f"1:{s}" for s in S.ANNO_SCALES]
    rec("Annotation scales", "Scale list = " + ", ".join(want), pf([s[0] for s in sc] == want),
        f"read back {[s[0] for s in sc]}; drawing units {[s[1] for s in sc]}")
    rec("Annotation scales", "No duplicate scales", pf(len({s[0] for s in sc}) == len(sc)))
    rec("Annotation scales", "Scale list accepted by AutoCAD", "UNVERIFIED",
        "Raw AcDbScale objects; confirm in SCALELISTEDIT (ASG-QA checks it)")

    tsd = doc.rootdict.get("ACAD_TABLESTYLE")
    have = [k for k, _ in tsd.items()]
    rec("Table styles", f"{len(S.TABLE_STYLES)} ASG-TABLE-* styles", "UNVERIFIED",
        f"ezdxf cannot write TABLESTYLE; created by ASG-SETUP in AutoCAD. In DXF now: {have or 'none'}")
    rec("Page setups", "Named page setups ASG-<layout>", "UNVERIFIED",
        "ezdxf cannot write named PLOTSETTINGS; ASG-SETUP copies them from the layouts")

    lnames = doc.layouts.names_in_taborder()
    want_l = ["Model"] + [s[0] for s in S.SHEETS]
    rec("Layouts", "Layouts: " + ", ".join(want_l[1:]), pf(lnames == want_l), str(lnames))
    for name, w, hh, paper, tb, vps in S.SHEETS:
        lay = doc.layouts.get(name)
        d = lay.dxf
        ok = (abs(d.paper_width - w) < .01 and abs(d.paper_height - hh) < .01
              and d.plot_configuration_file == S.PLOTTER
              and d.paper_size == f"{paper}_({w:.2f}_x_{hh:.2f}_MM)"
              and d.current_style_sheet == S.CTB_NAME and d.plot_layout_flags & 128   # print lw
              and d.plot_layout_flags & 4 and not d.plot_layout_flags & 1             # centred, no vp borders
              and d.plot_type == 5 and d.scale_denominator == 1 and d.scale_numerator == 1)
        rec("Layouts", f"{name}: {w:.0f}x{hh:.0f} mm, {S.PLOTTER}, layout area, 1:1, centred, "
                       f"{S.CTB_NAME}, plot lineweights", pf(bool(ok)), d.paper_size)
        ins = [e for e in lay if e.dxftype() == "INSERT"]
        tags = sorted(a.dxf.tag for a in ins[0].attribs) if ins else []
        okb = len(ins) == 1 and ins[0].dxf.name == tb and ins[0].dxf.xscale == 1 == ins[0].dxf.yscale
        rec("Title blocks", f"{name}: {tb} at scale 1 (not stretched), {len(tags)} attributes",
            pf(okb and tags == sorted(t for t, _p, _d in S.TB_ATTRIBUTES)))
        bx = [e for e in lay if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == "ASG-BORDER"]
        pts = [tuple(round(c, 3) for c in p[:2]) for p in bx[0].get_points()] if bx else []
        okx = pts == [(S.MARGIN_LEFT, S.MARGIN), (w - S.MARGIN, S.MARGIN), (w - S.MARGIN, hh - S.MARGIN),
                      (S.MARGIN_LEFT, hh - S.MARGIN)]
        rec("Sheets", f"{name}: border 20/10 mm margins inside {w:.0f}x{hh:.0f}", pf(okx))
        tbw, tbh = S.TB_SPECS[tb]["w"], None
        vps_ = [e for e in lay if e.dxftype() == "VIEWPORT" and e.dxf.id != 1]
        v = vps_[0] if vps_ else None
        okv = (v is not None and v.dxf.layer == "ASG-VIEWPORT" and v.dxf.flags & 16384
               and abs(v.dxf.view_height / v.dxf.height - vps) < 1e-6)
        rec("Viewports", f"{name}: 1 viewport, locked, ASG-VIEWPORT (no plot), 1:{vps}", pf(bool(okv)))
        # viewport must not overlap title block
        if v is not None:
            vx2 = v.dxf.center.x + v.dxf.width / 2
            vy1 = v.dxf.center.y - v.dxf.height / 2
            tb_x1 = w - S.MARGIN - tbw
            tb_y2 = S.MARGIN + S.TB_HEIGHTS_CHECK[tb]
            clear = vx2 <= tb_x1 or vy1 >= tb_y2
            rec("Viewports", f"{name}: viewport clear of title block", pf(clear))
    for tb, spec in S.TB_SPECS.items():
        blk = doc.blocks.get(tb)
        a = sorted(x.dxf.tag for x in blk.query("ATTDEF"))
        bb = ezdxf.bbox.extents(blk.query("LWPOLYLINE"))
        rec("Title blocks", f"{tb}: {len(a)} attribute definitions, {bb.size.x:.0f} x {bb.size.y:.1f} mm",
            pf(a == sorted(t for t, _p, _d in S.TB_ATTRIBUTES)))
    rec("Model space", "Model space empty", pf(len(doc.modelspace()) == 0), f"{len(doc.modelspace())} entities")
    xr = [b.name for b in doc.blocks if b.block and b.block.is_xref]
    rec("References", "No external references", pf(not xr))
    imgs = len(doc.objects.query("IMAGEDEF")) + len(doc.objects.query("UNDERLAYDEFINITION"))
    rec("References", "No image / PDF underlay dependencies", pf(imgs == 0))
    prox = [e.dxftype() for e in doc.entitydb.values() if e.dxftype().startswith("ACAD_PROXY")]
    rec("Interoperability", "No proxy entities", pf(not prox))
    ub = sorted(b.name for b in doc.blocks if not b.name.startswith("*") and not b.name.startswith("_"))
    rec("Scope", "Only title-block blocks (no product blocks / libraries)", pf(ub == sorted(S.TB_SPECS)), str(ub))
    fonts = sorted({s.dxf.font for s in doc.styles})
    rec("Interoperability", "Fonts referenced", pf(fonts == ["arial.ttf", "arialbd.ttf"]),
        f"{fonts} - Windows core fonts; AutoCAD substitutes via FONTALT if missing")
    return doc


def test_idempotency():
    before = BA.MASTER_DXF.read_bytes()
    S.inject_annotation_scales(BA.MASTER_DXF)
    again = BA.MASTER_DXF.read_bytes()
    rec("Automation", "Re-running scale injection adds nothing (idempotent)", pf(before == again))
    doc = S.new_document()
    doc2 = S.new_document()
    same = ([l.dxf.name for l in doc.layers] == [l.dxf.name for l in doc2.layers]
            and len(doc.dimstyles) == len(doc2.dimstyles) and doc.layouts.names() == doc2.layouts.names())
    rec("Automation", "Rebuild produces identical layer / style / layout sets", pf(same))
    S._layouts(doc, S.SHEETS)
    rec("Automation", "Re-running layout creation creates no duplicate layouts",
        pf(len(doc.layouts) == len(S.SHEETS) + 1))
    rec("Automation", "ASG_Setup.lsp idempotency", "UNVERIFIED", "Uses get-or-create; needs AutoCAD")


def test_templates():
    for name, *_ in S.SHEETS:
        p = BA.D_TEMPL / f"ASG_{name}.dxf"
        doc, aud = recover.readfile(p)
        ok = (not aud.has_errors and doc.layouts.names_in_taborder() == ["Model", name]
              and len(doc.modelspace()) == 0 and len([l for l in doc.layers if l.dxf.name.startswith("ASG-")]) == len(S.LAYERS))
        rec("Drawing templates", f"ASG_{name}.dxf: opens, single layout, full standard, empty model", pf(ok))


def test_ctb():
    p = BA.D_PLOT / S.CTB_NAME
    raw = p.read_bytes()
    hdr_ok = raw.startswith(b"PIAFILEVERSION_2.0,CTBVER1,compress\r\npmzlibcodec") and len(raw) > 60
    ctb = acadctb.load(str(p))
    blacks = all(ctb[a].color == (0, 0, 0) for a in range(1, 256) if a not in S.CTB_GREYS)
    greys = all(ctb[a].color == (g, g, g) for a, g in S.CTB_GREYS.items())
    lw = all(ctb[a].lineweight == 0 for a in range(1, 256))
    rec("Plot style", f"{S.CTB_NAME}: binary header (4-byte fields) + zlib body re-read", pf(hdr_ok),
        "ezdxf writer bug (8-byte fields on 64-bit Linux) worked around and verified")
    rec("Plot style", "CTB: ACI 1-249,255 -> black; 250-254 -> grey; object lineweight", pf(blacks and greys and lw))
    rec("Plot style", "CTB loads in AutoCAD Plot Style Manager", "UNVERIFIED", "README step 1")


def lineweight_test():
    """Render one line per lineweight class (ByLayer, real ASG layers) at 2400 dpi and measure."""
    doc = S.new_document(sheets=[])
    msp = doc.modelspace()
    for i, (lw, _l) in enumerate(S.LW_MATRIX):
        layer = next(l[1] for l in S.LAYERS if l[5] == lw and l[6])
        msp.add_line((5, 75 - i * 9), (95, 75 - i * 9), dxfattribs={"layer": layer})
    fig = plt.figure(figsize=(100 / 25.4, 80 / 25.4))
    ax = fig.add_axes([0, 0, 1, 1])
    cfg = Configuration(color_policy=ColorPolicy.BLACK, lineweight_policy=LineweightPolicy.ABSOLUTE,
                        background_policy=BackgroundPolicy.WHITE, min_lineweight=0.01)
    Frontend(RenderContext(doc), MatplotlibBackend(ax, adjust_figure=False), config=cfg).draw_layout(msp, finalize=True)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 80)
    ax.set_aspect("auto")
    p = D_QA / "QA_lineweight_matrix.png"
    fig.savefig(p, dpi=2400)
    plt.close(fig)
    img = plt.imread(p)[..., :3].mean(axis=2)
    H, W = img.shape
    pxmm = W / 100
    col = img[:, W // 2] < 0.5
    widths = []
    for i in range(len(S.LW_MATRIX)):
        row = int((80 - (75 - i * 9)) * pxmm)
        widths.append(col[row - 100:row + 100].sum() / pxmm)
    ratios = [w / lw for w, (lw, _l) in zip(widths, S.LW_MATRIX)]
    med = sorted(ratios)[len(ratios) // 2]
    ordered = all(a > b for a, b in zip(widths, widths[1:]))
    proportional = all(abs(r / med - 1) < 0.25 for r in ratios)
    rec("Lineweights", "Lineweight matrix (2400 dpi render): 8 classes distinct, ordered, proportional",
        "PARTIAL" if ordered and proportional else "FAIL",
        "nominal->rendered width ratio: " + ", ".join(f"{lw:.2f}:{r:.2f}" for r, (lw, _l) in zip(ratios, S.LW_MATRIX)) +
        ". Hierarchy confirmed; absolute plotted widths need AutoCAD PDF plot (renderer not mm-calibrated).")
    return widths


def test_qa_drawing():
    doc, aud = recover.readfile(BA.QA_DXF)
    rec("QA drawing", "Temporary QA drawing is a separate file", pf(BA.QA_DXF != BA.MASTER_DXF and not aud.has_errors),
        BA.QA_DXF.name)
    msp = doc.modelspace()
    dims = list(msp.query("DIMENSION"))
    meas = {}
    for d in dims:
        m = d.get_measurement()
        meas.setdefault(d.dxf.dimstyle, []).append(round(float(m if not hasattr(m, "x") else m.x), 4))
    rec("Dimensions", "Linear dims measure geometry (1200, 2400)",
        pf(sorted(meas.get("ASG-DIM-FIXED", [])) == [1200, 2400]), str(meas.get("ASG-DIM-FIXED")))
    rec("Dimensions", "Detail dim measures 50", pf(meas.get("ASG-DIM-DETAIL") == [50]))
    ang = meas.get("ASG-DIM-ANGULAR", [0])[0]
    rec("Dimensions", "Angular dim measures 30.00 deg", pf(abs(np.degrees(ang) - 30) < 1e-6 or abs(ang - 30) < 1e-6),
        f"{ang}")
    rec("Dimensions", "Radius / diameter dims present and measure R40 / D80",
        pf(meas.get("ASG-DIM-RADIUS") == [40] and meas.get("ASG-DIM-DIAMETER") == [80]),
        f"R {meas.get('ASG-DIM-RADIUS')} D {meas.get('ASG-DIM-DIAMETER')}")
    over = [d for d in dims if d.dxf.get("text", "<>") not in ("", "<>")]
    rec("Dimensions", "No overridden dimension text", pf(not over))
    texts = []
    for d in dims:
        for e in d.virtual_entities():
            if e.dxftype() in ("MTEXT", "TEXT"):
                texts.append(e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text)
    rec("Dimensions", "Rendered dim text precision: 1200 / 2400 (0 dp), 50.0 (detail 0.0), angles 0.00",
        pf("1200" in texts and "2400" in texts and "50.0" in texts and any(t.startswith("30.00") for t in texts)),
        str(texts))
    lay = doc.layouts.get("A3-LANDSCAPE")
    vps = [e for e in lay if e.dxftype() == "VIEWPORT" and e.dxf.id != 1]
    scales = sorted(round(v.dxf.view_height / v.dxf.height, 6) for v in vps)
    rec("Viewports", "QA viewports at 1:20 and 1:5", pf(scales == [5, 20]), str(scales))
    png = D_QA / "QA_render_A3.png"
    render(doc, lay, png, dpi=600, pdf=D_QA / "QA_render_A3.pdf")
    img = plt.imread(png)[..., :3].mean(axis=2)
    H, W = img.shape
    px_mm = W / 420.0
    widths = lineweight_test()
    # non-plot behaviour: construction diagonal + NOPLOT text inside 1:20 viewport
    # NOPLOT text at model (300,1200) in 1:20 viewport centred on (450,1100) at paper (75,148.5)
    px_, py_ = 75 + (300 - 450) / 20, 148.5 + (1200 - 1100) / 20
    r0, c0 = int((297 - py_) * H / 297), int(px_ * px_mm)
    probe = img[r0 - 60:r0 + 60, c0 - 120:c0 + 120]
    dark = probe.min() < 0.5
    rec("Plotting", "No-plot layers (ASG-CONSTRUCTION, ASG-NOPLOT) excluded from plot",
        "UNVERIFIED" if dark else "PARTIAL",
        "Plot flag is set and re-read (PASS above), but the ezdxf renderer ignores the layer plot flag - "
        "no-plot geometry appears in QA_render_A3. Confirm in AutoCAD plot preview." if dark else
        "probe area blank in render")
    rec("Plotting", "Viewport boundaries not plotted", "UNVERIFIED",
        "ASG-VIEWPORT plot flag off (PASS above); renderer never draws viewport frames, so not testable here")
    rec("Plotting", "PDF output via DWG To PDF.pc3 with CTB", "UNVERIFIED",
        "QA_render_A3.pdf produced by ezdxf/matplotlib, not by AutoCAD plot engine")
    rec("Annotation scales", "Annotative objects in viewports (text/dims/mleaders)", "UNVERIFIED",
        "ezdxf cannot create scale-context data; QA drawing uses fixed-scale overrides; test in AutoCAD")
    return widths


# ===========================================================================
def write_register():
    D_REG.mkdir(parents=True, exist_ok=True)
    header = ["Layer code", "Full layer name", "Functional group", "Purpose", "ACI colour", "Linetype",
              "Lineweight (mm)", "Plot status", "Description", "Recommended use", "Revision"]
    rows = [[c, n, g, d, a, lt, f"{lw:.2f}", "Plot" if p else "NO PLOT",
             f"{d}; ACI {a}, {lt}, {lw:.2f} mm", u, S.STANDARD_REV]
            for c, n, g, a, lt, lw, p, d, u in S.LAYERS]
    with open(D_REG / "ASG_Layer_Register.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([header] + rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Layer Register"
    ws.append(["ASG CAD MASTER STANDARD - LAYER REGISTER"])
    ws.append([f"{S.COMPANY} | {S.GROUP} | Standard {S.STANDARD_REV} | {TODAY}"])
    ws.append([])
    ws.append(header)
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    thin = Side(style="thin", color="A0A0A0")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3A3A3A")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=5, max_row=4 + len(rows)):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 9, 10))
        if row[7].value == "NO PLOT":
            row[7].font = Font(bold=True, color="C00000")
    for col, w in zip("ABCDEFGHIJK", (9, 24, 13, 26, 9, 12, 11, 10, 40, 46, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:K{4 + len(rows)}"
    ws2 = wb.create_sheet("Lineweight Matrix")
    ws2.append(["Lineweight (mm)", "Class", "Layers"])
    for lw, cls in S.LW_MATRIX:
        ws2.append([f"{lw:.2f}", cls, ", ".join(l[1] for l in S.LAYERS if l[5] == lw and l[6])])
    ws3 = wb.create_sheet("Colour Logic")
    for r in COLOUR_LOGIC:
        ws3.append(list(r))
    for w_ in (ws2, ws3):
        for c in w_[1]:
            c.font = Font(bold=True)
        w_.column_dimensions["B"].width = 50
        w_.column_dimensions["C"].width = 120
    wb.save(D_REG / "ASG_Layer_Register.xlsx")


COLOUR_LOGIC = [
    ("ACI", "Function", "Plotted (ASG_Monochrome.ctb)"),
    ("7", "Primary geometry, metals, text, sheet", "Black"),
    ("8", "Secondary geometry, hardware, seals, hidden, centre, grids", "Black (thinner weights)"),
    ("4", "Glass and mirror", "Black"),
    ("3", "Annotation: dimensions, leaders, levels, symbols", "Black"),
    ("1", "Revisions", "Black"),
    ("6", "Non-plot: construction, viewport, no-plot", "Not plotted (layer plot flag off)"),
    ("251 / 252 / 253", "Fills / hatch & insulation / reference", "Grey 36% / 52% / 68% darkness"),
]


# ===========================================================================
def pdf_doc(path, title, story_fn):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    def footer(c, d):
        c.saveState()
        c.setFont("Helvetica", 7)
        c.drawString(15 * mm, 8 * mm, f"{S.COMPANY}  |  {title}  |  Standard {S.STANDARD_REV}  |  {TODAY}")
        c.drawRightString(d.pagesize[0] - 15 * mm, 8 * mm, f"Page {d.page}")
        c.restoreState()

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                      topMargin=12 * mm, bottomMargin=14 * mm, title=title, author=S.COMPANY
                      ).build(story_fn(), onFirstPage=footer, onLaterPages=footer)


def styles():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    ss = getSampleStyleSheet()
    return {
        "H1": ParagraphStyle("h1", parent=ss["Heading1"], fontSize=15, spaceAfter=4),
        "H2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=3),
        "P": ParagraphStyle("p", parent=ss["BodyText"], fontSize=8.5, leading=11),
        "C": ParagraphStyle("c", parent=ss["BodyText"], fontSize=7.3, leading=8.8),
    }


def tbl(data, widths, st, status_col=None):
    from reportlab.lib import colors as rc
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle
    t = Table([[Paragraph(str(c), st["C"]) for c in r] for r in data],
              colWidths=[w * mm for w in widths], repeatRows=1)
    sty = [("GRID", (0, 0), (-1, -1), 0.3, rc.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
           ("BACKGROUND", (0, 0), (-1, 0), rc.HexColor("#d0d0d0")),
           ("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2)]
    if status_col is not None:
        fills = {"PASS": "#d9ead3", "PARTIAL": "#fff2cc", "UNVERIFIED": "#e6e6e6", "FAIL": "#f4cccc"}
        for i, r in enumerate(data[1:], 1):
            sty.append(("BACKGROUND", (status_col, i), (status_col, i), rc.HexColor(fills.get(r[status_col], "#ffffff"))))
    t.setStyle(TableStyle(sty))
    return t


def write_standards_pdf(previews):
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, Spacer
    st = styles()
    P, H2 = st["P"], st["H2"]

    def story():
        s = [Paragraph(S.COMPANY + "  |  " + S.GROUP, P),
             Paragraph("ASG CAD MASTER STANDARD - AutoCAD Environment for Aluminium, Glass &amp; Facade Engineering", st["H1"]),
             Paragraph(f"Standard {S.STANDARD_REV} | {TODAY} | Status: issued for internal review. "
                       "Conventions are informed by ISO 128 (line types / widths), ISO 5455 (scales), ISO 5457 "
                       "(sheet sizes / borders), ISO 3098 (lettering) and ISO 7200 (title block fields). "
                       "No formal compliance or certification is claimed - the ISO texts were not available "
                       "for verification.", P)]
        s += [Paragraph("1. Design philosophy", H2)] + [Paragraph(f"{i}. {t}", P) for i, t in enumerate(PHILOSOPHY, 1)]
        s += [Paragraph("2. Package structure", H2), tbl([["Folder", "Content"]] + PACKAGE, (55, 205), st)]
        s += [PageBreak(), Paragraph("3. Units and drawing-resident variables (stored in the DXF/DWG/DWT)", H2),
              tbl([["Variable", "Value", "Purpose"]] + [[v, str(x), p] for v, x, p in S.HEADER_VARS + S.HEADER_STYLE_VARS],
                  (35, 45, 180), st),
              Paragraph("Precision: ordinary dimensions 0 dp; fabrication details ASG-DIM-DETAIL 0.0; angles 0.00. "
                        "Rounding DIMRND 0 (no rounding beyond precision). Trailing zeros suppressed on 0-dp "
                        "styles; shown on DETAIL (0.0). INSUNITS = mm on template and every block prevents "
                        "unit conversion on insertion.", P)]
        s += [Paragraph("4. Machine / user-profile settings (NOT stored in any DWG/DWT)", H2),
              Paragraph("Stored in the AutoCAD profile / registry of each PC. A template cannot enforce them. "
                        "Recommended values; applied only by the opt-in command ASG-PROFILE.", P),
              tbl([["Variable", "Recommended", "Purpose"]] + [list(r) for r in S.PROFILE_VARS], (45, 45, 170), st)]
        s += [PageBreak(), Paragraph("5. Layer dictionary", H2),
              Paragraph("Naming: ASG-&lt;FUNCTION&gt;[-&lt;QUALIFIER&gt;], full words, no ambiguous abbreviations. "
                        "Codes: G = geometry, T = technical, A = annotation, S = sheet. All objects ByLayer.", P),
              tbl([["Code", "Layer", "ACI", "Linetype", "LW", "Plot", "Recommended use"]] +
                  [[c, n, a, lt, f"{lw:.2f}", "Yes" if p else "NO", u] for c, n, _g, a, lt, lw, p, _d, u in S.LAYERS],
                  (12, 45, 10, 20, 11, 11, 151), st),
              Spacer(1, 4), tbl([list(r) for r in COLOUR_LOGIC], (30, 120, 110), st)]
        s += [PageBreak(), Paragraph("6. Lineweight matrix and linetypes", H2),
              tbl([["Lineweight", "Class"]] + [[f"{a:.2f} mm", b] for a, b in S.LW_MATRIX], (30, 230), st),
              Paragraph("Steps between classes are at least 0.04 mm (ISO 128 ratio ~1.4 between main weights). "
                        "LTSCALE 1 + PSLTSCALE 1: dash lengths are paper mm in every viewport; MSLTSCALE 1 "
                        "(profile) applies the annotation scale in model space.", P),
              tbl([["Linetype", "Pattern"]] + [[k, v[1]] for k, v in S.LINETYPES.items()] + [["Continuous", "Solid"]],
                  (40, 220), st)]
        s += [Paragraph("7. Text system - Arial", H2),
              tbl([["Style", "Font", "Annotative", "Use"]] +
                  [[n, f, "Yes" if a else "No (paper space)", u] for n, f, _b, a, u in S.TEXT_STYLES], (40, 25, 30, 165), st),
              Spacer(1, 3),
              tbl([["Use", "Plotted height", "Style"]] + [[a, f"{b} mm", c] for a, b, c in S.TEXT_HEIGHTS], (60, 30, 170), st),
              Paragraph("Font availability: Arial / Arial Bold are Windows core fonts present on every company PC "
                        "and embed in PDF. If missing, AutoCAD substitutes FONTALT (set FONTALT = arial.ttf). "
                        "Heights are set in paper mm; annotative styles scale them automatically.", P)]
        s += [PageBreak(), Paragraph("8. Dimension system", H2),
              tbl([["Style", "Annotative", "Use"]] + [[n, "Yes" if a else "No", u] for n, a, _o, u in S.DIM_STYLES],
                  (40, 20, 200), st),
              tbl([["Parameter", "Value"]] + DIM_PARAMS, (60, 200), st),
              Paragraph("Rule: dimensions always report measured geometry (DIMLFAC 1, DIMASSOC 2). Overriding "
                        "dimension text is prohibited; fix the geometry instead. No further fixed-scale styles "
                        "are created - ASG-DIM-FIXED in paper space covers non-annotative workflows.", P)]
        s += [Paragraph("9. Multileader system", H2),
              tbl([["Style", "Annotative", "Arrow", "Text", "Frame", "Use"]] +
                  [[n, "Yes" if a else "No", (ar or "closed filled") + f" {sz}", f"{ts} {ch}", "Yes" if fr else "No", u]
                   for n, a, ar, sz, ts, ch, fr, u in S.MLEADER_STYLES], (35, 18, 30, 40, 13, 124), st),
              Paragraph("All: straight leader, 2 points max, landing gap 1.0, dogleg 4.0, attachment middle of top "
                        "line, ByLayer colour. Background mask: per MText (Properties &gt; Text &gt; Background "
                        "mask, border offset 1.2, use drawing background) where notes cross hatch.", P)]
        s += [PageBreak(), Paragraph("10. Annotative scale governance", H2),
              Paragraph("Controlled list: " + ", ".join(f"1:{x}" for x in S.ANNO_SCALES) +
                        ". Additional scales only with CAD custodian approval (e.g. 1:200 for site plans).", P),
              tbl([["Object", "Rule"]] + ANNO_RULES, (40, 220), st),
              Paragraph("<b>Workflow:</b> 1) Set the viewport scale and lock the viewport. 2) Double-click into the "
                        "viewport - CANNOSCALE follows it. 3) Add annotation with ASG-*-ANNO styles. 4) Need the same "
                        "note in another scale? Select > right-click > Annotative Object Scale > Add (OBJECTSCALE). "
                        "5) Keep ANNOAUTOSCALE 0 and ANNOALLVISIBLE 0 to avoid duplicates. 6) Imported drawings with "
                        "extra scales: -SCALELISTEDIT &gt; Reset, then AUDIT.", P)]
        s += [Paragraph("11. Tables and schedules", H2),
              tbl([["Style", "Title", "Header", "Data", "Use"]] +
                  [[n, "suppressed" if ts_ else f"{t[0]} {t[1]}", "suppressed" if hs else f"{h[0]} {h[1]}",
                    f"{d[0]} {d[1]}", u] for n, t, h, d, ts_, hs, u in S.TABLE_STYLES], (40, 30, 30, 30, 130), st),
              Paragraph("Common: ASG-TEXT-SHEET, cell margins H 1.5 / V 1.0, flow top-down, inside grid 0.18, outline "
                        "0.35, header fill ACI 254 (light grey), no data fill. Row height = text + 2 x margin "
                        "(min 6 mm at 2.5). Numbers right/centre, descriptions left. Place tables in paper space on "
                        "ASG-TABLE. Created by ASG-SETUP.", P)]
        s += [PageBreak(), Paragraph("12. Layouts, page setups and title block", H2),
              tbl([["Layout / page setup", "Paper (canonical name)", "Orient.", "Title block", "Default viewport"]] +
                  [[f"{n} / ASG-{n}", f"{p}_({w:.2f}_x_{h:.2f}_MM)", "L" if w > h else "P", tb, f"1:{v} locked"]
                   for n, w, h, p, tb, v in S.SHEETS], (45, 85, 14, 25, 91), st),
              Paragraph("All page setups: DWG To PDF.pc3, plot area Layout, scale 1:1 (paper units mm), plot offset "
                        "centred, ASG_Monochrome.ctb, plot object lineweights ON, lineweights not scaled, plot "
                        "transparency OFF, paper space last, viewport borders on no-plot layer ASG-VIEWPORT. "
                        "Border margins 20 mm left (filing) / 10 mm others (ISO 5457 practice). Viewport practice: "
                        "every viewport on ASG-VIEWPORT, scale set from the controlled list, Display Locked = Yes; "
                        "unlock only to change the scale.", P),
              Paragraph("Title block: separate A4 (180 mm), A3 (180 mm, larger rows) and A1 (250 mm) designs, "
                        "inserted at scale 1. Hierarchy bottom-up: drawing number / revision / sheet; scale / date / "
                        "status; signatures; drawing title; client / consultant; project / location; company "
                        "heading (text, no logo invented). Attributes: " +
                        ", ".join(t for t, _p, _d in S.TB_ATTRIBUTES) + ".", P)]
        s += [Paragraph("13. Hatch, block and attribute standards", H2),
              tbl([["Material", "Pattern", "Scale", "Angle", "Layer", "Note"]] + [list(r) for r in S.HATCH_STANDARD],
                  (40, 20, 15, 12, 30, 143), st),
              Paragraph("Model-space hatch scale = table scale x drawing scale. Blocks: name ASG-&lt;CATEGORY&gt;-"
                        "&lt;ITEM&gt;-&lt;VARIANT&gt;, geometry on layer 0 ByBlock/ByLayer, base point at a datum "
                        "(fixing point / bottom-left / centreline), INSUNITS mm, scale uniformly, real-size items "
                        "non-annotative, symbols annotative. Manufacturer data only from verified supplier sources.", P),
              tbl([["Attribute tag", "Meaning"]] + [list(r) for r in S.ATTRIBUTE_STANDARD], (40, 220), st)]
        s += [PageBreak(), Paragraph("14. Plotting / PDF standard and interoperability", H2),
              tbl([["Item", "Standard"]] + PLOT_STD, (45, 215), st),
              Paragraph("15. Resource separation and portability checklist", H2),
              tbl([["Resource", "Where it lives", "Deploy"]] + RESOURCES, (45, 80, 135), st)]
        s += [PageBreak(), Paragraph("Appendix - layout previews (rendered from the saved master DXF)", H2)]
        for p in previews:
            s += [Paragraph(p.stem, P), Image(str(p), width=165 * mm, height=118 * mm, kind="proportional")]
        return s

    pdf_doc(D_DOC / "ASG_CAD_Standards.pdf", "ASG CAD Standards", story)


def write_validation(widths):
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph
    st = styles()
    counts = {k: sum(1 for r in R if r[2] == k) for k in ("PASS", "PARTIAL", "UNVERIFIED", "FAIL")}
    sha = hashlib.sha256(BA.MASTER_DXF.read_bytes()).hexdigest()[:16]

    def story():
        s = [Paragraph("ASG TEMPLATE VALIDATION REPORT", st["H1"]),
             Paragraph(f"{S.COMPANY} | master: 01_Master_Template/ASG_Master_Template.dxf (sha256 {sha}...) | "
                       f"ezdxf {ezdxf.__version__}, LibreOffice 24.2 (2nd reader), Python | {TODAY}", st["P"]),
             Paragraph("<b>Summary:</b> " + " | ".join(f"{k} {v}" for k, v in counts.items()) +
                       ". AutoCAD was NOT available: nothing below was tested inside AutoCAD. PASS = tested by "
                       "re-opening the saved file; PARTIAL = tested with a non-AutoCAD renderer only; "
                       "UNVERIFIED = needs AutoCAD (steps in README); FAIL = test failed or deliverable missing.", st["P"]),
             Paragraph("Test procedure", st["H2"])] + [Paragraph(f"{i}. {t}", st["P"]) for i, t in enumerate(PROCEDURE, 1)]
        s += [Paragraph("Results", st["H2"]),
              tbl([["Area", "Test", "Status", "Evidence"]] + [list(r) for r in R], (30, 95, 20, 115), st, status_col=2)]
        s += [PageBreak(), Paragraph("QA render - temporary QA drawing (A3, viewports 1:20 and 1:5, lineweight matrix, "
                                     "text heights). Not part of the master template.", st["P"]),
              Image(str(D_QA / "QA_render_A3.png"), width=240 * mm, height=170 * mm, kind="proportional")]
        return s

    pdf_doc(D_QA / "ASG_Template_Validation_Report.pdf", "Template Validation Report", story)
    lines = [f"ASG TEMPLATE VALIDATION REPORT - {TODAY}",
             "Summary: " + " | ".join(f"{k} {v}" for k, v in counts.items()),
             "AutoCAD not available - nothing tested inside AutoCAD.", ""]
    lines += [f"{s:<11} [{a}] {t}" + (f"  -- {e}" if e else "") for a, t, s, e in R]
    (D_QA / "ASG_Template_Validation_Report.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return counts


PHILOSOPHY = [
    "Model geometry is drawn full size, 1:1, in millimetres.",
    "Geometry lives in model space; annotation is annotative in model space or fixed in paper space; sheets live in layouts.",
    "Layers are meaningful, coded, documented; object properties are ByLayer.",
    "Blocks have predictable names, datum base points and mm insertion units.",
    "Dimensions always report true geometry - no text overrides.",
    "Annotation is sized by plotted paper height, never by guesswork in model space.",
    "Layouts use genuine ISO paper sizes and locked, controlled viewports.",
    "Plots are monochrome with a fixed lineweight hierarchy - predictable on every printer and PDF.",
    "Proprietary manufacturer data is never invented.",
    "Easy to learn, hard to misuse: few styles, clear names, one setup command, one QA command.",
]
PACKAGE = [
    ["01_Master_Template", "ASG_Master_Template.dxf (+ .dwg/.dwt after AutoCAD step), ASG_Setup.lsp"],
    ["02_Drawing_Templates", "Single-layout templates per sheet size (DXF)"],
    ["03_Page_Setups", "ASG_Page_Setups.csv - page setup definitions (created in AutoCAD by ASG-PAGESETUPS)"],
    ["04_Plotting_Standards", "ASG_Monochrome.ctb"],
    ["05_CAD_Standards_Documentation", "ASG_CAD_Standards.pdf (this document)"],
    ["06_Layer_Register", "ASG_Layer_Register.xlsx / .csv"],
    ["07_QA_Validation", "Validation report PDF/TXT, temporary QA drawing, QA renders"],
    ["08_Automation_Source", "asg_standard.py, build_all.py, qa_validate.py, ASG_Setup.lsp"],
]
DIM_PARAMS = [
    ["Text", "ASG-TEXT-ANNO 2.5, above line, aligned with line, gap 1.0"],
    ["Arrowheads", "Closed filled 2.5; centre mark 2.5 (DIAMETER), none (RADIUS)"],
    ["Extension lines", "Offset 1.5 from origin, extend 1.25 beyond dim line"],
    ["Spacing", "Baseline / continue spacing 7.0"],
    ["Fit", "Best fit (DIMATFIT 3), text moved beside dim line, force line inside"],
    ["Linear units", "Decimal mm, 0 dp (DETAIL 0.0), '.' separator, rounding 0, trailing zeros suppressed (DETAIL shown)"],
    ["Angular units", "Decimal degrees 0.00"],
    ["Tolerance", "Off; symmetrical +/- precision 0.0 height 0.7 switched on per dimension on DETAIL"],
    ["Alternate units", "Off"],
    ["Measurement scale", "DIMLFAC 1.0 - true geometry"],
    ["Scale", "Annotative styles: DIMSCALE from annotation scale; FIXED: DIMSCALE 1 for paper space"],
]
ANNO_RULES = [
    ["Model-space text", "Annotative ASG-TEXT-ANNO / -NOTE / -SMALL"],
    ["Paper-space text", "Non-annotative ASG-TEXT-TITLE / -HEADING / -SHEET at true plotted height"],
    ["Dimensions", "ASG-DIM-ANNO family in model space; ASG-DIM-FIXED for paper-space dims"],
    ["Multileaders", "ASG-ML-ANNO / -DETAIL / -REFERENCE in model space; ASG-ML-NOTE in paper space"],
    ["Symbols", "Section / detail / level / grid bubbles annotative"],
    ["Real-size blocks", "Never annotative (hardware, profiles, fixings)"],
    ["Hatches", "Material hatches non-annotative; symbolic fills may be annotative"],
]
PLOT_STD = [
    ["Device", "DWG To PDF.pc3 (AutoCAD standard) - no custom PC3 dependency"],
    ["Paper", "ISO full bleed A4 / A3 / A1 - exact sheet size; border gives margins"],
    ["Area / scale / offset", "Layout, 1:1, centred"],
    ["Plot style", "ASG_Monochrome.ctb (black; 250-254 grey); fallback monochrome.ctb"],
    ["Lineweights", "Plot object lineweights ON, scale lineweights OFF"],
    ["Transparency / shading", "Plot transparency OFF; shaded viewports 'As displayed', 300 dpi"],
    ["PDF options", "Vector 1200 dpi, raster 300 dpi, TrueType as text, layer info OFF for client issue, hyperlinks off"],
    ["Pre-issue", "Freeze/check no-plot layers, AUDIT, PURGE, plot preview, open the PDF and check scale bar / dims"],
    ["Compatibility", "Save AutoCAD 2018 DWG; DXF 2018 for third parties; no proxy objects, no custom objects"],
    ["Xrefs / PDF import", "Relative paths, overlay, ASG-REFERENCE; PDFIMPORT geometry moved to ASG layers"],
]
RESOURCES = [
    ["Layers, styles, scales, layouts, title blocks", "Inside DWT (drawing-resident)", "Copy DWT to shared Templates folder"],
    ["Named page setups", "Inside DWT after ASG-SETUP", "Travel with DWT"],
    ["Plot style ASG_Monochrome.ctb", "Plot Styles folder (external)", "Copy to each PC / shared support path"],
    ["DWG To PDF.pc3", "AutoCAD Plotters folder (ships with AutoCAD)", "None"],
    ["Fonts arial.ttf / arialbd.ttf", "Windows Fonts (external)", "Present on Windows; check on Mac/viewers"],
    ["Linetypes", "Embedded in DWT (acadiso values)", "None"],
    ["Hatch patterns", "acadiso.pat (ships with AutoCAD)", "None"],
    ["OSMODE, grips, autosave, background ...", "User profile / registry (external)", "Optional ASG-PROFILE per user"],
    ["QNEW default template", "Options > Files (profile)", "Set per PC to ASG_Master_Template.dwt"],
]
PROCEDURE = [
    "Inventory CAD software (AutoCAD, accoreconsole, ODA, LibreDWG, QCAD, BricsCAD).",
    "Build all files from asg_standard.py (build_all.py).",
    "Re-open master with ezdxf recover + audit, strict parse, and LibreOffice Draw as an independent reader.",
    "Compare every header variable, layer, linetype, text/dim/mleader/mline style, scale and layout with the definitions.",
    "Check page setup fields, borders, title-block insert scale and attributes, viewport layer/lock/scale/position.",
    "Check model space empty, no xrefs/underlays/proxies, only title-block blocks, fonts.",
    "Idempotency: re-run scale injection and layout creation; rebuild and compare.",
    "Re-open each single-size template and the CTB (binary header + all 255 entries).",
    "Build a separate temporary QA drawing: dims of known geometry, viewports 1:20 / 1:5, no-plot objects, lineweight matrix, text heights.",
    "Check measured dimension values and rendered text precision; render QA layout at 600 dpi and measure line widths; probe no-plot and viewport-edge areas.",
    "In AutoCAD (pending): ASG-SETUP, ASG-QA, plot QA layout to PDF, check annotative scaling, then SAVEAS DWG/DWT.",
]


def previews_of_master():
    doc = ezdxf.readfile(BA.MASTER_DXF)
    out = []
    pv = D_DOC / "previews"
    pv.mkdir(parents=True, exist_ok=True)
    for lay in doc.layouts:
        if lay.name == "Model":
            continue
        p = pv / f"{lay.name}.png"
        render(doc, lay, p, dpi=110)
        out.append(p)
    return out


def main():
    S.TB_HEIGHTS.clear()
    S.new_document()  # populate title block heights for geometry checks
    S.TB_HEIGHTS_CHECK = dict(S.TB_HEIGHTS)
    test_environment()
    test_master(BA.MASTER_DXF)
    test_idempotency()
    test_templates()
    test_ctb()
    widths = test_qa_drawing()
    write_register()
    rec("Deliverables", "ASG_Layer_Register.xlsx written and re-opened",
        pf(__import__("openpyxl").load_workbook(D_REG / "ASG_Layer_Register.xlsx")["Layer Register"].max_row == 4 + len(S.LAYERS)))
    pv = previews_of_master()
    write_standards_pdf(pv)
    rec("Deliverables", "ASG_CAD_Standards.pdf written", pf((D_DOC / "ASG_CAD_Standards.pdf").stat().st_size > 10000))
    counts = write_validation(widths)
    print(counts)
    for a, t, s, e in R:
        if s != "PASS":
            print(f"{s:<11} {a}: {t} -- {e}")


if __name__ == "__main__":
    main()
