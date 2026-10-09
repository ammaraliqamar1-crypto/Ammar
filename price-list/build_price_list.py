"""Builds ASG_Material_Price_List.xlsx - category-wise material & rate database
for Aluminium, Glass, MS & SS works. Run: python3 build_price_list.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

OUT = "ASG_Material_Price_List.xlsx"
FONT = "Arial"
NAVY, GOLD, GREY = "1F3864", "C9A227", "F2F4F7"
SPARE_ROWS = 15  # blank rows reserved at the end of each category for new items

# ---------------------------------------------------------------- settings
# (name, label, value, unit, note)
SETTINGS = [
    ("VAT", "VAT (UAE)", 0.05, "%", "Federal VAT rate"),
    ("OHP", "Overhead & Profit (default)", 0.20, "%", "Applied to Net Cost to get Selling Rate"),
    ("AL_MILL", "Aluminium extrusion - mill finish", 15.50, "AED/kg", "6063-T5/T6, local extruder"),
    ("AL_PC", "Add for powder coating", 3.00, "AED/kg", "Qualicoat approved, standard RAL"),
    ("AL_ANOD", "Add for anodizing", 5.00, "AED/kg", "15-20 micron"),
    ("AL_PVDF", "Add for PVDF coating", 8.50, "AED/kg", "2/3-coat Kynar 500"),
    ("AL_WOOD", "Add for wood-grain finish", 6.50, "AED/kg", "Sublimation over powder coat"),
    ("AL_TB", "Add for thermal-break (polyamide)", 6.00, "AED/kg", "Crimped polyamide strips"),
    ("MS_KG", "MS structural sections", 3.60, "AED/kg", "SHS/RHS/angle/channel/beam, black"),
    ("MS_PLATE", "MS plates & sheets", 3.40, "AED/kg", "S275/S355, cut-to-size extra"),
    ("GI_KG", "GI sheets / pre-galvanized", 4.20, "AED/kg", "Z275"),
    ("SS304", "SS 304 tubes / sheets (hairline)", 17.00, "AED/kg", "Grade 304, #4 hairline"),
    ("SS316", "SS 316 tubes / sheets (hairline)", 24.00, "AED/kg", "Grade 316/316L, marine/external"),
    ("SS_MIRROR", "Add for SS mirror finish (#8)", 3.00, "AED/kg", "Over hairline rate"),
]

# Rate formulas below may use the names above, e.g. "=ROUND(MS_KG*4.25,2)".
AL_PCR = "(AL_MILL+AL_PC)"


def al(w):  # aluminium powder-coated per lm from weight kg/m
    return f"=ROUND({w}*{AL_PCR},2)"


def ms(w):
    return f"=ROUND({w}*MS_KG,2)"


def msp(w):
    return f"=ROUND({w}*MS_PLATE,2)"


def ss(w, g="SS304"):
    return f"=ROUND({w}*{g},2)"


# ---------------------------------------------------------------- data
# Each category: (sheet, prefix, title, scope, tab colour, default wastage, items)
# item: (sub-category, description, specification/size, finish, brand/origin, unit, rate[, wastage])
CATS = []


def cat(sheet, prefix, title, scope, colour, wastage, items):
    CATS.append(dict(sheet=sheet, prefix=prefix, title=title, scope=scope,
                     colour=colour, wastage=wastage, items=items))


cat("01 Aluminium", "AL", "Aluminium Profiles & Systems",
    "Extrusions by kg, standard sections, window/door/curtain wall/pergola profiles", "8EA9DB", 0.05, [
        ("Extrusions (per kg)", "Aluminium extrusion", "Alloy 6063-T5/T6", "Mill finish", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL"),
        ("Extrusions (per kg)", "Aluminium extrusion", "Alloy 6063-T5/T6", "Powder coated (RAL)", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL+AL_PC"),
        ("Extrusions (per kg)", "Aluminium extrusion", "Alloy 6063-T5/T6", "Anodized 15-20 micron", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL+AL_ANOD"),
        ("Extrusions (per kg)", "Aluminium extrusion", "Alloy 6063-T5/T6", "PVDF 2/3-coat", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL+AL_PVDF"),
        ("Extrusions (per kg)", "Aluminium extrusion", "Alloy 6063-T5/T6", "Wood-grain finish", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL+AL_PC+AL_WOOD"),
        ("Extrusions (per kg)", "Thermal-break aluminium extrusion", "With polyamide strips", "Powder coated", "Gulf Extrusions / Equiv.", "kg", "=AL_MILL+AL_PC+AL_TB"),
        ("Standard Sections", "Aluminium square tube", "25 x 25 x 1.5 mm (0.38 kg/m)", "Powder coated", "Local", "lm", al(0.38)),
        ("Standard Sections", "Aluminium square tube", "40 x 40 x 2.0 mm (0.82 kg/m)", "Powder coated", "Local", "lm", al(0.82)),
        ("Standard Sections", "Aluminium square tube", "50 x 50 x 2.0 mm (1.04 kg/m)", "Powder coated", "Local", "lm", al(1.04)),
        ("Standard Sections", "Aluminium rectangular tube", "100 x 50 x 2.5 mm (1.96 kg/m)", "Powder coated", "Local", "lm", al(1.96)),
        ("Standard Sections", "Aluminium rectangular tube", "100 x 50 x 3.0 mm (2.33 kg/m)", "Powder coated", "Local", "lm", al(2.33)),
        ("Standard Sections", "Aluminium rectangular tube", "150 x 50 x 3.0 mm (3.14 kg/m)", "Powder coated", "Local", "lm", al(3.14)),
        ("Standard Sections", "Aluminium rectangular tube", "200 x 50 x 3.0 mm (3.95 kg/m)", "Powder coated", "Local", "lm", al(3.95)),
        ("Standard Sections", "Aluminium equal angle", "25 x 25 x 2.0 mm (0.26 kg/m)", "Powder coated", "Local", "lm", al(0.26)),
        ("Standard Sections", "Aluminium equal angle", "50 x 50 x 3.0 mm (0.79 kg/m)", "Powder coated", "Local", "lm", al(0.79)),
        ("Standard Sections", "Aluminium flat bar", "50 x 5 mm (0.68 kg/m)", "Powder coated", "Local", "lm", al(0.68)),
        ("Standard Sections", "Aluminium flat bar", "100 x 6 mm (1.62 kg/m)", "Powder coated", "Local", "lm", al(1.62)),
        ("Standard Sections", "Aluminium U-channel for glass", "For 10-12 mm glass (0.45 kg/m)", "Powder coated / anodized", "Local", "lm", al(0.45)),
        ("Standard Sections", "Aluminium heavy U-channel (frameless partition)", "40 x 40 mm (0.90 kg/m)", "Anodized / PC", "Local", "lm", al(0.90)),
        ("Sliding Window System", "Sliding window outer frame - 2 track", "Non-thermal (approx. 1.25 kg/m)", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(1.25)),
        ("Sliding Window System", "Sliding window outer frame - 3 track", "Non-thermal (approx. 1.75 kg/m)", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(1.75)),
        ("Sliding Window System", "Sliding sash profile", "Approx. 0.95 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(0.95)),
        ("Sliding Window System", "Sliding interlock profile", "Approx. 0.75 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(0.75)),
        ("Sliding Window System", "Thermal-break sliding frame", "Approx. 2.40 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", f"=ROUND(2.4*(AL_MILL+AL_PC+AL_TB),2)"),
        ("Sliding Window System", "Lift & slide door frame (TB)", "Approx. 3.80 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", f"=ROUND(3.8*(AL_MILL+AL_PC+AL_TB),2)"),
        ("Casement / Hinged System", "Casement window frame", "Approx. 0.95 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(0.95)),
        ("Casement / Hinged System", "Casement window sash", "Approx. 1.05 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(1.05)),
        ("Casement / Hinged System", "Window mullion / transom", "Approx. 1.30 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(1.30)),
        ("Casement / Hinged System", "Thermal-break casement frame", "Approx. 1.50 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", f"=ROUND(1.5*(AL_MILL+AL_PC+AL_TB),2)"),
        ("Casement / Hinged System", "Hinged door frame", "Approx. 1.60 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(1.60)),
        ("Casement / Hinged System", "Hinged door leaf (stile)", "Approx. 2.10 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(2.10)),
        ("Casement / Hinged System", "Door bottom rail (deep)", "Approx. 2.80 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(2.80)),
        ("Casement / Hinged System", "Glazing bead", "Approx. 0.25 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(0.25)),
        ("Curtain Wall System", "Curtain wall mullion - 150 mm deep", "Stick system (approx. 3.80 kg/m)", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(3.80)),
        ("Curtain Wall System", "Curtain wall mullion - 180 mm deep", "Stick system (approx. 4.60 kg/m)", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(4.60)),
        ("Curtain Wall System", "Curtain wall mullion - 200 mm deep", "Stick system (approx. 5.40 kg/m)", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(5.40)),
        ("Curtain Wall System", "Curtain wall transom - 120 mm", "Approx. 2.40 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(2.40)),
        ("Curtain Wall System", "Pressure plate", "Approx. 0.55 kg/m", "Mill finish", "Gulf Extrusions / Equiv.", "lm", "=ROUND(0.55*AL_MILL,2)"),
        ("Curtain Wall System", "Snap-on cover cap", "Approx. 0.45 kg/m", "Powder coated / PVDF", "Gulf Extrusions / Equiv.", "lm", al(0.45)),
        ("Curtain Wall System", "Unitized frame profile (male/female)", "Approx. 4.80 kg/m", "PVDF", "Gulf Extrusions / Equiv.", "lm", "=ROUND(4.8*(AL_MILL+AL_PVDF),2)"),
        ("Curtain Wall System", "Structural glazing (SSG) carrier frame", "Approx. 0.85 kg/m", "Mill finish", "Gulf Extrusions / Equiv.", "lm", "=ROUND(0.85*AL_MILL,2)"),
        ("Pergola / Louver / Skylight", "Pergola main beam", "150 x 50 x 3 mm (3.14 kg/m)", "Powder coated / wood-grain", "Local", "lm", al(3.14)),
        ("Pergola / Louver / Skylight", "Pergola rafter / slat", "100 x 50 x 2.5 mm (1.96 kg/m)", "Powder coated / wood-grain", "Local", "lm", al(1.96)),
        ("Pergola / Louver / Skylight", "Bioclimatic pergola louver blade", "Approx. 2.60 kg/m", "Powder coated", "Local / Imported", "lm", al(2.60)),
        ("Pergola / Louver / Skylight", "Louver blade - Z type 100 mm", "Approx. 0.85 kg/m", "Powder coated", "Local", "lm", al(0.85)),
        ("Pergola / Louver / Skylight", "Louver blade - aerofoil 200 mm", "Approx. 2.30 kg/m", "Powder coated", "Local", "lm", al(2.30)),
        ("Pergola / Louver / Skylight", "Skylight rafter profile", "Approx. 3.20 kg/m", "Powder coated", "Gulf Extrusions / Equiv.", "lm", al(3.20)),
        ("Balustrade / Handrail", "Glass balustrade base shoe profile", "Heavy duty, for 12-21.52 mm glass (6.8 kg/m)", "Anodized / PC", "Local / Imported", "lm", al(6.80)),
        ("Balustrade / Handrail", "Aluminium handrail - oval / round", "Approx. 0.90 kg/m", "Powder coated", "Local", "lm", al(0.90)),
        ("Balustrade / Handrail", "Glass top cap (U-cap) profile", "For 12-17.52 mm glass (0.55 kg/m)", "Anodized / PC", "Local", "lm", al(0.55)),
        ("Sheet & Panel", "Aluminium flat sheet", "2.0 mm, 1250 x 2500", "Mill finish", "Local / Imported", "m2", 48.0),
        ("Sheet & Panel", "Aluminium flat sheet", "3.0 mm, 1250 x 2500", "Mill finish", "Local / Imported", "m2", 70.0),
        ("Sheet & Panel", "Aluminium chequered plate", "3.0 mm tread", "Mill finish", "Local / Imported", "m2", 85.0),
    ])

cat("02 Glass", "GL", "Glass - Float, Tempered, Laminated & DGU",
    "Clear, tinted, reflective, tempered, laminated, insulated, low-E, low-iron, fire-rated", "9BC2E6", 0.05, [
        ("Annealed Float", "Clear float glass", "4 mm", "Clear", "Emirates Float / Guardian", "m2", 30.0),
        ("Annealed Float", "Clear float glass", "6 mm", "Clear", "Emirates Float / Guardian", "m2", 40.0),
        ("Annealed Float", "Clear float glass", "8 mm", "Clear", "Emirates Float / Guardian", "m2", 55.0),
        ("Annealed Float", "Clear float glass", "10 mm", "Clear", "Emirates Float / Guardian", "m2", 70.0),
        ("Annealed Float", "Clear float glass", "12 mm", "Clear", "Emirates Float / Guardian", "m2", 88.0),
        ("Annealed Float", "Tinted float glass", "6 mm", "Grey / Bronze / Green", "Emirates Float / Guardian", "m2", 55.0),
        ("Annealed Float", "Reflective glass (hard coat)", "6 mm", "Blue / Grey / Silver", "Emirates Float / Guardian", "m2", 75.0),
        ("Tempered", "Clear tempered glass", "6 mm, edges polished", "Clear", "Local processor", "m2", 75.0),
        ("Tempered", "Clear tempered glass", "8 mm, edges polished", "Clear", "Local processor", "m2", 95.0),
        ("Tempered", "Clear tempered glass", "10 mm, edges polished", "Clear", "Local processor", "m2", 120.0),
        ("Tempered", "Clear tempered glass", "12 mm, edges polished", "Clear", "Local processor", "m2", 150.0),
        ("Tempered", "Clear tempered glass", "15 mm, edges polished", "Clear", "Local processor", "m2", 260.0),
        ("Tempered", "Clear tempered glass", "19 mm, edges polished", "Clear", "Local processor", "m2", 420.0),
        ("Tempered", "Tinted tempered glass", "10 mm", "Grey / Bronze", "Local processor", "m2", 145.0),
        ("Tempered", "Low-iron (extra clear) tempered glass", "10 mm", "Ultra clear", "Local processor", "m2", 210.0),
        ("Tempered", "Low-iron (extra clear) tempered glass", "12 mm", "Ultra clear", "Local processor", "m2", 260.0),
        ("Tempered", "Frosted / acid-etched tempered glass", "10 mm", "Satin", "Local processor", "m2", 170.0),
        ("Tempered", "Frosted / acid-etched tempered glass", "12 mm", "Satin", "Local processor", "m2", 205.0),
        ("Tempered", "Fluted / reeded tempered glass", "8 mm", "Fluted", "Imported", "m2", 230.0),
        ("Tempered", "Heat-soak test (HST) - extra", "Any thickness", "-", "Local processor", "m2", 25.0),
        ("Tempered", "Heat-strengthened glass - extra over tempered", "Any thickness", "-", "Local processor", "m2", 15.0),
        ("Laminated", "Clear laminated glass", "6.38 mm (3+0.38 PVB+3)", "Clear", "Local processor", "m2", 95.0),
        ("Laminated", "Clear laminated glass", "8.76 mm (4+0.76+4)", "Clear", "Local processor", "m2", 125.0),
        ("Laminated", "Clear laminated glass", "10.76 mm (5+0.76+5)", "Clear", "Local processor", "m2", 150.0),
        ("Laminated", "Clear laminated glass", "12.76 mm (6+0.76+6)", "Clear", "Local processor", "m2", 175.0),
        ("Laminated", "Tempered laminated glass", "12.76 mm (6T+0.76 PVB+6T)", "Clear", "Local processor", "m2", 260.0),
        ("Laminated", "Tempered laminated glass", "17.52 mm (8T+1.52 PVB+8T)", "Clear", "Local processor", "m2", 360.0),
        ("Laminated", "Tempered laminated glass", "21.52 mm (10T+1.52 PVB+10T)", "Clear", "Local processor", "m2", 450.0),
        ("Laminated", "Tempered laminated glass - SGP interlayer", "21.52 mm (10T+1.52 SGP+10T)", "Clear", "Local processor", "m2", 650.0),
        ("Laminated", "Laminated glass - frosted / milky PVB", "10.76 mm", "Milky white", "Local processor", "m2", 175.0),
        ("Insulated (DGU)", "Double glazed unit - clear", "6 + 12 air + 6 mm", "Clear", "Local processor", "m2", 180.0),
        ("Insulated (DGU)", "Double glazed unit - tinted", "6 tinted + 12 + 6 clear", "Tinted", "Local processor", "m2", 210.0),
        ("Insulated (DGU)", "Double glazed unit - reflective", "6 reflective + 12 + 6 clear", "Reflective", "Local processor", "m2", 230.0),
        ("Insulated (DGU)", "Double glazed unit - Low-E", "6 HT Low-E + 12 + 6 HT clear", "Low-E soft coat", "Guardian SunGuard / Equiv.", "m2", 280.0),
        ("Insulated (DGU)", "Double glazed unit - Low-E with laminated inner", "6 HT Low-E + 12 + 6.38 lam", "Low-E soft coat", "Guardian SunGuard / Equiv.", "m2", 360.0),
        ("Insulated (DGU)", "Double glazed unit - Low-E, argon filled", "8 HT Low-E + 16 argon + 8 HT", "Low-E soft coat", "Guardian SunGuard / Equiv.", "m2", 380.0),
        ("Insulated (DGU)", "Triple glazed unit", "6 + 12 + 6 + 12 + 6 mm", "Clear / Low-E", "Local processor", "m2", 420.0),
        ("Insulated (DGU)", "Structural silicone (SSG) DGU - extra", "Structural edge seal", "-", "Local processor", "m2", 35.0),
        ("Decorative & Special", "Back-painted (lacquered) tempered glass", "6 mm", "RAL colour", "Local processor", "m2", 140.0),
        ("Decorative & Special", "Back-painted (lacquered) tempered glass", "10 mm", "RAL colour", "Local processor", "m2", 185.0),
        ("Decorative & Special", "Ceramic frit / spandrel glass", "6 mm HT, opaque", "RAL colour", "Local processor", "m2", 150.0),
        ("Decorative & Special", "Digital printed glass", "10 mm tempered", "Custom print", "Local processor", "m2", 380.0),
        ("Decorative & Special", "Smart (switchable PDLC) glass", "12.76 mm laminated", "Switchable", "Imported", "m2", 1400.0),
        ("Decorative & Special", "Wired glass", "6.8 mm", "Clear", "Imported", "m2", 110.0),
        ("Decorative & Special", "Anti-slip floor glass (laminated)", "21.52 mm, fritted top", "Anti-slip frit", "Local processor", "m2", 750.0),
        ("Fire-Rated", "Fire-rated glass - E30 (integrity)", "Approx. 6-8 mm", "Clear", "Pyrobel / Pyrostop / Equiv.", "m2", 900.0),
        ("Fire-Rated", "Fire-rated glass - EI30 (insulated)", "Approx. 16-18 mm", "Clear", "Pyrobel / Pyrostop / Equiv.", "m2", 1600.0),
        ("Fire-Rated", "Fire-rated glass - EI60 (insulated)", "Approx. 23-27 mm", "Clear", "Pyrobel / Pyrostop / Equiv.", "m2", 2400.0),
        ("Fire-Rated", "Fire-rated glass - EI90/120", "Approx. 37-50 mm", "Clear", "Pyrobel / Pyrostop / Equiv.", "m2", 3800.0),
        ("Polycarbonate & Acrylic", "Polycarbonate multiwall sheet", "10 mm twin wall", "Clear / Opal / Bronze", "Imported", "m2", 55.0),
        ("Polycarbonate & Acrylic", "Polycarbonate multiwall sheet", "16 mm multiwall", "Clear / Opal", "Imported", "m2", 85.0),
        ("Polycarbonate & Acrylic", "Polycarbonate solid sheet", "6 mm", "Clear / Bronze", "Imported", "m2", 160.0),
        ("Polycarbonate & Acrylic", "Polycarbonate solid sheet", "10 mm", "Clear", "Imported", "m2", 260.0),
        ("Polycarbonate & Acrylic", "Acrylic (PMMA) sheet", "5 mm", "Clear / Opal", "Imported", "m2", 140.0),
    ])

cat("03 Glass Processing", "GP", "Glass Processing & Fabrication Charges",
    "Edge work, holes, cut-outs, bevelling, sand-blasting, film, cutting & handling", "BDD7EE", 0.0, [
        ("Edge Work", "Edge grinding (arris / seamed)", "Up to 12 mm", "-", "Local processor", "lm", 4.0),
        ("Edge Work", "Flat polished edge", "Up to 12 mm", "-", "Local processor", "lm", 7.0),
        ("Edge Work", "Flat polished edge", "15-19 mm", "-", "Local processor", "lm", 12.0),
        ("Edge Work", "Bevelled edge - 10 mm width", "Up to 6 mm mirror / glass", "-", "Local processor", "lm", 12.0),
        ("Edge Work", "Bevelled edge - 25 mm width", "Up to 6 mm mirror / glass", "-", "Local processor", "lm", 20.0),
        ("Edge Work", "Pencil / round edge", "Up to 12 mm", "-", "Local processor", "lm", 10.0),
        ("Edge Work", "Mitre edge (45 degree)", "Up to 12 mm", "-", "Local processor", "lm", 18.0),
        ("Holes & Cut-outs", "Hole drilling", "Up to 12 mm glass, dia up to 20 mm", "-", "Local processor", "nos", 10.0),
        ("Holes & Cut-outs", "Hole drilling", "Up to 12 mm glass, dia 21-50 mm", "-", "Local processor", "nos", 18.0),
        ("Holes & Cut-outs", "Countersunk hole", "For spider / routel fittings", "-", "Local processor", "nos", 30.0),
        ("Holes & Cut-outs", "Hinge / patch notch cut-out", "Up to 12 mm glass", "-", "Local processor", "nos", 30.0),
        ("Holes & Cut-outs", "Lock / handle cut-out", "Up to 12 mm glass", "-", "Local processor", "nos", 40.0),
        ("Holes & Cut-outs", "Socket / switch cut-out", "Up to 6 mm (splash back)", "-", "Local processor", "nos", 25.0),
        ("Holes & Cut-outs", "Corner radius / shaped cutting", "Per corner", "-", "Local processor", "nos", 10.0),
        ("Surface Treatment", "Sand-blasting / frosting", "Full sheet", "Frosted", "Local processor", "m2", 35.0),
        ("Surface Treatment", "Sand-blasted design / logo", "As per design", "Frosted", "Local processor", "m2", 75.0),
        ("Surface Treatment", "Frosted vinyl film (manifestation)", "Installed", "Frosted", "3M / Equiv.", "m2", 55.0),
        ("Surface Treatment", "Manifestation strip (safety band)", "100 mm wide, installed", "Frosted", "3M / Equiv.", "lm", 12.0),
        ("Surface Treatment", "Solar / safety window film", "Installed", "Tinted / clear", "3M / Llumar / Equiv.", "m2", 85.0),
        ("Surface Treatment", "Self-cleaning / anti-fingerprint coating", "Applied", "-", "Imported", "m2", 60.0),
        ("Handling", "Glass cutting charges (customer glass)", "Per cut", "-", "Local processor", "nos", 3.0),
        ("Handling", "Oversize glass surcharge", "Above 2.4 x 3.6 m", "-", "Local processor", "m2", 45.0),
        ("Handling", "Packing / crating (export / long distance)", "Wooden crate", "-", "Local processor", "m2", 15.0),
    ])

cat("04 Mirrors", "MR", "Mirrors",
    "Silver, coloured, safety-backed, bevelled, LED & framed mirrors", "DDEBF7", 0.05, [
        ("Silver Mirror", "Silver mirror", "4 mm", "Clear silver", "Guardian / Saint-Gobain / Equiv.", "m2", 45.0),
        ("Silver Mirror", "Silver mirror", "5 mm", "Clear silver", "Guardian / Saint-Gobain / Equiv.", "m2", 55.0),
        ("Silver Mirror", "Silver mirror", "6 mm", "Clear silver", "Guardian / Saint-Gobain / Equiv.", "m2", 65.0),
        ("Silver Mirror", "Copper-free / lead-free silver mirror", "6 mm", "Clear silver", "Imported", "m2", 80.0),
        ("Silver Mirror", "Safety-backed mirror (vinyl film)", "6 mm, Cat. 2 film", "Clear silver", "Local processor", "m2", 85.0),
        ("Silver Mirror", "Tempered mirror", "6 mm", "Clear silver", "Local processor", "m2", 120.0),
        ("Coloured Mirror", "Tinted mirror", "5 mm", "Bronze / Grey / Gold", "Imported", "m2", 75.0),
        ("Coloured Mirror", "Antique mirror", "5 mm", "Antique finish", "Imported", "m2", 220.0),
        ("Coloured Mirror", "Two-way (one-way) mirror", "6 mm", "Semi-reflective", "Imported", "m2", 260.0),
        ("Finished Mirror", "Mirror with 25 mm bevelled edge", "5 mm, fixed with SS clips", "Clear silver", "Local", "m2", 140.0),
        ("Finished Mirror", "Mirror with polished edge & adhesive fixing", "6 mm", "Clear silver", "Local", "m2", 120.0),
        ("Finished Mirror", "LED back-lit mirror with demister", "800 x 1000 mm", "Touch switch", "Imported", "nos", 750.0),
        ("Finished Mirror", "Framed mirror - aluminium frame", "Per m2", "Powder coated", "Local", "m2", 240.0),
        ("Finished Mirror", "Gym / full-height wall mirror", "6 mm on MDF/plywood backing", "Clear silver", "Local", "m2", 160.0),
        ("Mirror Accessories", "Mirror adhesive (neutral silicone)", "310 ml", "-", "Dow / Sika / Equiv.", "nos", 25.0),
        ("Mirror Accessories", "SS mirror clip / J-channel", "Set", "SS 304", "Local", "set", 12.0),
        ("Mirror Accessories", "Aluminium J-channel / mirror channel", "Per lm", "Anodized", "Local", "lm", 15.0),
    ])

cat("05 Door & Window HW", "HW", "Hardware - Doors & Windows",
    "Floor springs, closers, locks, handles, hinges, rollers, friction stays, multipoint locks", "F4B084", 0.0, [
        ("Floor Springs & Closers", "Floor spring", "Up to 150 kg, EN 1-4", "SS cover", "dormakaba BTS 75V / Equiv.", "nos", 850.0),
        ("Floor Springs & Closers", "Floor spring", "Up to 120 kg", "SS cover", "Local / China", "nos", 300.0),
        ("Floor Springs & Closers", "Overhead door closer", "EN 2-4, with arm", "Silver", "dormakaba TS 83 / Equiv.", "nos", 380.0),
        ("Floor Springs & Closers", "Overhead door closer - slide arm", "EN 1-4", "Silver", "GEZE TS 4000 / Equiv.", "nos", 450.0),
        ("Floor Springs & Closers", "Concealed door closer", "EN 2-3", "-", "dormakaba ITS 96 / Equiv.", "nos", 950.0),
        ("Floor Springs & Closers", "Overhead door closer", "EN 2-4", "Silver", "Local / China", "nos", 120.0),
        ("Floor Springs & Closers", "Door stopper - floor mounted", "-", "SS satin", "Local", "nos", 25.0),
        ("Locks & Cylinders", "Mortise lock - narrow stile (aluminium door)", "Backset 25/30 mm", "SS", "Giesse / CISA / Equiv.", "nos", 120.0),
        ("Locks & Cylinders", "Euro profile cylinder", "70 mm, key/key", "Satin nickel", "dormakaba / Equiv.", "nos", 75.0),
        ("Locks & Cylinders", "Euro profile cylinder", "70 mm, key/thumb-turn", "Satin nickel", "dormakaba / Equiv.", "nos", 90.0),
        ("Locks & Cylinders", "Multipoint lock (hinged door)", "2-4 locking points", "-", "Giesse / Savio / Equiv.", "nos", 180.0),
        ("Locks & Cylinders", "Sliding door hook lock", "With cylinder", "-", "Giesse / Equiv.", "nos", 85.0),
        ("Locks & Cylinders", "Sliding window crescent lock", "-", "Silver / Black", "Local / China", "nos", 15.0),
        ("Locks & Cylinders", "Electromagnetic lock", "600 lbs, with bracket", "-", "Local / China", "nos", 220.0),
        ("Locks & Cylinders", "Panic bar (push bar)", "Single door", "SS / Silver", "dormakaba / Equiv.", "nos", 650.0),
        ("Handles", "Pull handle - H type", "600 mm, dia 32 mm", "SS 304 satin", "Local", "pair", 120.0),
        ("Handles", "Pull handle - H type", "1200 mm, dia 32 mm", "SS 304 satin", "Local", "pair", 220.0),
        ("Handles", "Pull handle - H type", "1800 mm, dia 38 mm", "SS 316 satin", "Imported", "pair", 450.0),
        ("Handles", "Lever handle set on rose", "-", "SS satin", "Hoppe / Equiv.", "set", 140.0),
        ("Handles", "Lever handle set on rose", "-", "SS satin", "Local / China", "set", 45.0),
        ("Handles", "Window handle (casement)", "-", "Silver / Black", "Giesse / Savio / Equiv.", "nos", 25.0),
        ("Handles", "Sliding door flush handle", "-", "Silver / Black", "Local", "nos", 18.0),
        ("Hinges & Pivots", "Aluminium door hinge (3D adjustable)", "Up to 100 kg", "Powder coated", "Giesse / Savio / Equiv.", "nos", 45.0),
        ("Hinges & Pivots", "Butt hinge SS", "100 x 75 x 3 mm", "SS 304", "Local", "nos", 15.0),
        ("Hinges & Pivots", "Concealed hinge (aluminium door)", "Up to 80 kg", "-", "Imported", "nos", 120.0),
        ("Hinges & Pivots", "Pivot set (top & bottom)", "Up to 250 kg", "SS", "FritsJurgens / Equiv.", "set", 1800.0),
        ("Window Hardware", "Friction stay", "300-400 mm", "SS 304", "Local / Imported", "pair", 45.0),
        ("Window Hardware", "Friction stay - heavy duty", "500-600 mm", "SS 316", "Imported", "pair", 95.0),
        ("Window Hardware", "Tilt & turn hardware set", "Per sash", "-", "Roto / Siegenia / Equiv.", "set", 350.0),
        ("Window Hardware", "Restrictor / limiter stay", "100 mm opening", "SS", "Local / Imported", "nos", 25.0),
        ("Window Hardware", "Sliding window roller", "Single, up to 40 kg", "-", "Local / China", "nos", 12.0),
        ("Window Hardware", "Sliding door roller - tandem", "Up to 150 kg", "SS bearing", "Giesse / Equiv.", "nos", 65.0),
        ("Window Hardware", "Lift & slide hardware set", "Per sash, up to 300 kg", "-", "Siegenia / Roto / Equiv.", "set", 1800.0),
        ("Window Hardware", "Mosquito / insect screen (sliding)", "Aluminium frame + mesh", "PC", "Local", "m2", 85.0),
        ("Window Hardware", "Pleated insect screen", "Per m2", "PC", "Imported", "m2", 260.0),
        ("Weather Seals", "Brush seal (wool pile)", "Per lm", "Grey / Black", "Local", "lm", 1.5),
        ("Weather Seals", "Automatic door bottom seal", "Per door", "-", "Local / Imported", "nos", 120.0),
        ("Weather Seals", "Aluminium threshold with seal", "Per lm", "Anodized", "Local", "lm", 45.0),
    ])

cat("06 Glass Fittings HW", "GF", "Glass Fittings - Partition, Shower & Spider",
    "Patch fittings, shower hinges, clamps, spider / routel, sliding glass systems", "F8CBAD", 0.0, [
        ("Patch Fittings (Glass Doors)", "Bottom patch fitting", "For 10-12 mm glass", "SS satin", "dormakaba PT 10 / Equiv.", "nos", 220.0),
        ("Patch Fittings (Glass Doors)", "Top patch fitting", "For 10-12 mm glass", "SS satin", "dormakaba PT 20 / Equiv.", "nos", 220.0),
        ("Patch Fittings (Glass Doors)", "Top pivot / over-panel fitting", "For 10-12 mm glass", "SS satin", "dormakaba PT 30 / Equiv.", "nos", 260.0),
        ("Patch Fittings (Glass Doors)", "Corner lock patch with strike", "For 10-12 mm glass", "SS satin", "dormakaba PT 40 / Equiv.", "set", 320.0),
        ("Patch Fittings (Glass Doors)", "Patch fitting set (top, bottom, lock, strike)", "Complete door set", "SS satin", "Local / China", "set", 180.0),
        ("Patch Fittings (Glass Doors)", "Transom / side-panel connector", "Glass-to-glass", "SS satin", "Local / Imported", "nos", 85.0),
        ("Shower Fittings", "Shower hinge 90 deg - glass to wall", "For 8-10 mm glass", "SS 304 chrome / satin", "Local / China", "nos", 85.0),
        ("Shower Fittings", "Shower hinge 180 deg - glass to glass", "For 8-10 mm glass", "SS 304 chrome / satin", "Local / China", "nos", 95.0),
        ("Shower Fittings", "Shower hinge 135 deg - glass to glass", "For 8-10 mm glass", "SS 304 chrome / satin", "Local / China", "nos", 95.0),
        ("Shower Fittings", "Shower hinge - premium", "For 8-10 mm glass", "Chrome / gold / black", "dormakaba / Häfele / Equiv.", "nos", 350.0),
        ("Shower Fittings", "Shower door knob / pull handle", "300 mm back-to-back", "SS chrome / satin", "Local", "pair", 60.0),
        ("Shower Fittings", "Glass-to-wall clamp (L / flat)", "For 8-10 mm glass", "SS chrome", "Local / China", "nos", 30.0),
        ("Shower Fittings", "Glass-to-glass clamp (90/180 deg)", "For 8-10 mm glass", "SS chrome", "Local / China", "nos", 35.0),
        ("Shower Fittings", "Support bar / stabilizer bar with fittings", "Up to 1000 mm", "SS chrome", "Local / China", "set", 140.0),
        ("Shower Fittings", "Sliding shower kit (barn type)", "Up to 1500 mm, 8-10 mm glass", "SS satin / black", "Local / China", "set", 450.0),
        ("Shower Fittings", "Shower seal / drip strip (PVC)", "Per lm", "Clear", "Local", "lm", 12.0),
        ("Shower Fittings", "Magnetic seal profile", "Per lm", "Clear", "Local", "lm", 22.0),
        ("Partition Hardware", "Sliding glass door track system", "Up to 80 kg, 2 m track", "Anodized", "dormakaba MUTO / Equiv.", "set", 1800.0),
        ("Partition Hardware", "Sliding glass door track system", "Up to 80 kg, 2 m track", "Anodized", "Local / China", "set", 650.0),
        ("Partition Hardware", "Soft-close damper (sliding)", "Per damper", "-", "Imported", "nos", 120.0),
        ("Partition Hardware", "Floor guide for sliding glass", "-", "SS", "Local", "nos", 25.0),
        ("Partition Hardware", "Stacking / folding glass partition hardware", "Per panel", "SS", "dormakaba HSW / Equiv.", "panel", 1600.0),
        ("Partition Hardware", "Glass door lock (centre / floor)", "With cylinder", "SS satin", "Local / Imported", "nos", 120.0),
        ("Spider & Point-Fixed", "Spider fitting - 4 arm", "SS 316, 200-250 mm c/c", "SS 316 satin", "Imported", "nos", 280.0),
        ("Spider & Point-Fixed", "Spider fitting - 2 arm", "SS 316", "SS 316 satin", "Imported", "nos", 190.0),
        ("Spider & Point-Fixed", "Spider fitting - 1 arm", "SS 316", "SS 316 satin", "Imported", "nos", 140.0),
        ("Spider & Point-Fixed", "Routel / articulated bolt", "Countersunk, for 12-21.52 mm glass", "SS 316", "Imported", "nos", 120.0),
        ("Spider & Point-Fixed", "Glass fin bracket", "SS 316", "SS 316", "Imported", "nos", 220.0),
        ("Spider & Point-Fixed", "Glass canopy rod / tie-rod set", "SS 316, adjustable", "SS 316", "Imported", "set", 650.0),
        ("Spider & Point-Fixed", "Glass canopy bracket / gripper", "SS 316", "SS 316", "Imported", "nos", 180.0),
    ])

cat("07 Balustrade HW", "BH", "Balustrade & Handrail Components",
    "SS posts, glass clamps, spigots, standoffs, handrail brackets & end caps", "FFE699", 0.0, [
        ("Posts", "SS balustrade post with glass clamps", "Dia 50.8 mm, 1000-1100 mm H", "SS 304 satin", "Local / China", "nos", 260.0),
        ("Posts", "SS balustrade post with glass clamps", "Dia 50.8 mm, 1000-1100 mm H", "SS 316 satin", "Local / China", "nos", 340.0),
        ("Posts", "SS square post with glass clamps", "40 x 40 mm, 1000-1100 mm H", "SS 304 satin", "Local / China", "nos", 280.0),
        ("Posts", "SS post with wire / rod infill holes", "Dia 50.8 mm, 5 holes", "SS 316 satin", "Local", "nos", 320.0),
        ("Glass Clamps & Spigots", "Glass clamp - round / flat back", "For 8-12 mm glass", "SS 304", "Local / China", "nos", 35.0),
        ("Glass Clamps & Spigots", "Glass clamp - round / flat back", "For 8-12 mm glass", "SS 316", "Local / China", "nos", 50.0),
        ("Glass Clamps & Spigots", "Glass spigot - core drilled", "For 12-17.52 mm glass", "SS 2205 duplex", "Imported", "nos", 260.0),
        ("Glass Clamps & Spigots", "Glass spigot - base plate", "For 12-17.52 mm glass", "SS 316", "Local / China", "nos", 180.0),
        ("Glass Clamps & Spigots", "Glass standoff / point fixing", "Dia 50 mm, for 12-17.52 mm glass", "SS 316", "Local / China", "nos", 75.0),
        ("Glass Clamps & Spigots", "Base shoe wedge / gasket kit", "Per lm", "EPDM / PVC", "Local / Imported", "lm", 40.0),
        ("Glass Clamps & Spigots", "Base shoe end cap", "-", "Aluminium / SS cladding", "Local", "nos", 45.0),
        ("Glass Clamps & Spigots", "Base shoe SS cladding / cover", "Per lm", "SS 304 satin", "Local", "lm", 120.0),
        ("Handrail Fittings", "Handrail bracket - wall mounted", "For dia 42-50 mm", "SS 304", "Local / China", "nos", 35.0),
        ("Handrail Fittings", "Handrail bracket - wall mounted", "For dia 42-50 mm", "SS 316", "Local / China", "nos", 50.0),
        ("Handrail Fittings", "Glass-mounted handrail bracket", "For 10-17.52 mm glass", "SS 316", "Local / China", "nos", 65.0),
        ("Handrail Fittings", "Handrail end cap / bend / elbow", "Dia 50.8 mm", "SS 304", "Local", "nos", 25.0),
        ("Handrail Fittings", "Handrail slotted top rail (glass cap)", "Dia 50.8 mm, 24 mm slot", "SS 304 satin", "Local / China", "lm", 110.0),
        ("Handrail Fittings", "Handrail slotted top rail (glass cap)", "Dia 50.8 mm, 24 mm slot", "SS 316 satin", "Local / China", "lm", 150.0),
        ("Wire & Rod Infill", "SS wire rope", "Dia 4 mm, 7x7", "SS 316", "Imported", "lm", 8.0),
        ("Wire & Rod Infill", "Wire rope tensioner / end fitting", "Set", "SS 316", "Imported", "set", 75.0),
        ("Wire & Rod Infill", "SS cross bar / rod", "Dia 12 mm", "SS 304 satin", "Local", "lm", 18.0),
        ("Wire & Rod Infill", "Cross bar holder", "Dia 12 mm", "SS 304", "Local", "nos", 10.0),
        ("Wire & Rod Infill", "SS perforated / mesh infill panel", "1.5 mm", "SS 304", "Local", "m2", 380.0),
    ])

cat("08 Sealants & Fixings", "SF", "Sealants, Gaskets, Fixings & Consumables",
    "Silicones, foams, EPDM, tapes, anchors, screws, brackets, shims, tools & consumables", "C6E0B4", 0.10, [
        ("Silicone Sealants", "Structural silicone sealant (2-part)", "Per kg (base + curing agent)", "Black", "Dow DOWSIL 993 / Equiv.", "kg", 45.0),
        ("Silicone Sealants", "Structural silicone sealant (1-part)", "600 ml sausage", "Black", "Dow DOWSIL 795 / 895", "nos", 34.0),
        ("Silicone Sealants", "Weatherproofing silicone sealant", "600 ml sausage", "Colour as req.", "Dow DOWSIL 791 / 991", "nos", 28.0),
        ("Silicone Sealants", "Weatherproofing silicone sealant", "300 ml cartridge", "Colour as req.", "Dow / Sika / GE", "nos", 16.0),
        ("Silicone Sealants", "Neutral cure general silicone", "300 ml cartridge", "Clear / White / Black", "GE / Local", "nos", 10.0),
        ("Silicone Sealants", "Anti-fungal sanitary silicone (shower)", "300 ml cartridge", "Clear / White", "Dow / Sika / Equiv.", "nos", 18.0),
        ("Silicone Sealants", "Insulating glass (IG) secondary sealant", "Per kg", "Black", "Dow DOWSIL 3362 / Equiv.", "kg", 38.0),
        ("Other Sealants & Adhesives", "Polyurethane (PU) sealant", "600 ml sausage", "Grey", "Sika Sikaflex / Equiv.", "nos", 30.0),
        ("Other Sealants & Adhesives", "PU expanding foam", "750 ml", "-", "Local / Imported", "nos", 18.0),
        ("Other Sealants & Adhesives", "Fire-rated sealant (intumescent / acrylic)", "310 ml", "White / Grey", "Hilti / Sika / Equiv.", "nos", 45.0),
        ("Other Sealants & Adhesives", "Fire-rated PU foam", "750 ml", "-", "Hilti / Equiv.", "nos", 55.0),
        ("Other Sealants & Adhesives", "Chemical anchor resin", "330-410 ml", "-", "Hilti HIT-RE / Fischer / Equiv.", "nos", 75.0),
        ("Other Sealants & Adhesives", "Primer / cleaner (silicone)", "1 litre", "-", "Dow / Sika", "ltr", 65.0),
        ("Other Sealants & Adhesives", "Masking tape", "48 mm x 50 m", "-", "Local", "roll", 6.0),
        ("Gaskets & Tapes", "EPDM glazing gasket (wedge)", "Per lm", "Black", "Local / Imported", "lm", 2.5),
        ("Gaskets & Tapes", "EPDM curtain wall gasket", "Per lm", "Black", "Local / Imported", "lm", 4.0),
        ("Gaskets & Tapes", "Silicone gasket (high performance)", "Per lm", "Black", "Imported", "lm", 7.0),
        ("Gaskets & Tapes", "Backer rod (closed-cell PE)", "Dia 10-25 mm", "-", "Local", "lm", 0.5),
        ("Gaskets & Tapes", "Double-sided structural glazing tape", "Per lm", "-", "3M VHB / Norton / Equiv.", "lm", 6.0),
        ("Gaskets & Tapes", "Foam spacer tape (SSG)", "Per lm", "-", "Norton / Equiv.", "lm", 2.0),
        ("Gaskets & Tapes", "Butyl / EPDM waterproofing membrane", "150-300 mm wide", "-", "Local / Imported", "lm", 12.0),
        ("Gaskets & Tapes", "Setting block (EPDM / silicone)", "100 x 25 mm", "-", "Local", "nos", 1.0),
        ("Gaskets & Tapes", "Glazing shims / packers (assorted)", "Packet", "-", "Local", "pkt", 15.0),
        ("Anchors & Fasteners", "Mechanical anchor bolt (wedge)", "M10 x 100", "Zinc plated", "Hilti HST3 / Fischer / Equiv.", "nos", 8.0),
        ("Anchors & Fasteners", "Mechanical anchor bolt (wedge)", "M12 x 120", "Zinc plated", "Hilti HST3 / Fischer / Equiv.", "nos", 11.0),
        ("Anchors & Fasteners", "SS anchor bolt", "M10 x 100", "SS 316", "Hilti / Fischer / Equiv.", "nos", 18.0),
        ("Anchors & Fasteners", "Chemical anchor stud with nut & washer", "M12 x 160", "HDG / SS 316", "Local / Imported", "nos", 14.0),
        ("Anchors & Fasteners", "Frame fixing (nylon plug + screw)", "10 x 100 mm", "-", "Fischer / Local", "nos", 1.2),
        ("Anchors & Fasteners", "Self-drilling screw", "4.8 x 19 / 25 mm", "SS 410 / Zinc", "Local", "box", 35.0),
        ("Anchors & Fasteners", "SS self-tapping screw (assorted)", "Box of 500", "SS 304", "Local", "box", 45.0),
        ("Anchors & Fasteners", "SS bolt, nut & washer set", "M10 x 50", "SS 304", "Local", "set", 3.0),
        ("Anchors & Fasteners", "SS bolt, nut & washer set", "M12 x 60", "SS 316", "Local", "set", 6.5),
        ("Anchors & Fasteners", "Pop rivet (aluminium / SS)", "Box of 1000", "-", "Local", "box", 40.0),
        ("Brackets & Embeds", "Curtain wall bracket (aluminium)", "Approx. 1.5 kg, with serrated plate", "Mill finish", "Local", "nos", 45.0),
        ("Brackets & Embeds", "MS galvanized bracket / cleat", "Approx. 2 kg", "HDG", "Local fabrication", "nos", 30.0),
        ("Brackets & Embeds", "Cast-in channel (Halfen type) with T-bolts", "Per lm", "HDG", "Halfen / Jordahl / Equiv.", "lm", 140.0),
        ("Brackets & Embeds", "Corner cleat / joint connector (aluminium)", "Per corner", "Mill finish", "Local", "nos", 6.0),
        ("Brackets & Embeds", "Isolation pad / separator (PVC/neoprene)", "Per nos", "-", "Local", "nos", 1.5),
        ("Consumables", "Cutting disc", "14 inch (aluminium / steel)", "-", "Local / Imported", "nos", 18.0),
        ("Consumables", "Grinding / flap disc", "4.5 inch", "-", "Local / Imported", "nos", 6.0),
        ("Consumables", "Welding electrode (E6013)", "2.5 / 3.2 mm, 5 kg pkt", "-", "ESAB / Equiv.", "pkt", 55.0),
        ("Consumables", "TIG filler rod SS 308L / 316L", "Per kg", "-", "ESAB / Equiv.", "kg", 65.0),
        ("Consumables", "Argon gas cylinder refill", "Per cylinder", "-", "Local", "nos", 120.0),
        ("Consumables", "Drill bits (HSS / masonry assorted)", "Set", "-", "Bosch / Equiv.", "set", 85.0),
        ("Consumables", "Protective film for aluminium", "Per m2", "-", "Local", "m2", 1.2),
        ("Consumables", "Cleaning materials (IPA, cloth, etc.)", "Lump sum per project", "-", "Local", "ls", 250.0),
    ])

cat("09 Cladding & ACP", "CL", "Cladding, ACP & Facade Panels",
    "ACP, solid aluminium, perforated, louvered and HPL panels", "A9D08E", 0.10, [
        ("ACP Sheets", "Aluminium composite panel (ACP)", "4 mm, 0.3 mm skin, non-FR", "PVDF", "Alubond / Alucobond / Equiv.", "m2", 55.0),
        ("ACP Sheets", "Aluminium composite panel - fire rated FR (B1)", "4 mm, 0.5 mm skin", "PVDF", "Alubond / Alucobond / Equiv.", "m2", 75.0),
        ("ACP Sheets", "Aluminium composite panel - A2 non-combustible", "4 mm, 0.5 mm skin", "PVDF", "Alubond / Alucobond / Equiv.", "m2", 110.0),
        ("ACP Sheets", "ACP - special finish (wood / stone / mirror / metallic)", "4 mm, FR", "Special", "Alubond / Equiv.", "m2", 130.0),
        ("Solid Aluminium Panels", "Solid aluminium panel", "2.0 mm", "PVDF", "Local fabrication", "m2", 150.0),
        ("Solid Aluminium Panels", "Solid aluminium panel", "3.0 mm", "PVDF", "Local fabrication", "m2", 190.0),
        ("Solid Aluminium Panels", "Perforated aluminium panel", "3.0 mm, pattern as approved", "PVDF", "Local fabrication", "m2", 260.0),
        ("Solid Aluminium Panels", "CNC laser-cut decorative screen", "3.0 mm, custom design", "PVDF", "Local fabrication", "m2", 380.0),
        ("Solid Aluminium Panels", "Aluminium honeycomb panel", "20-25 mm", "PVDF", "Imported", "m2", 450.0),
        ("Solid Aluminium Panels", "Aluminium corrugated / profiled sheet", "0.7 mm", "PC / PVDF", "Local", "m2", 65.0),
        ("Sub-Frame & Accessories", "ACP sub-frame (aluminium angles / tubes)", "Per m2 of cladding", "Mill finish", "Local", "m2", 45.0),
        ("Sub-Frame & Accessories", "GI sub-frame (Omega / Z channel)", "Per m2 of cladding", "Galvanized", "Local", "m2", 30.0),
        ("Sub-Frame & Accessories", "Rockwool insulation behind cladding", "50 mm, 100 kg/m3, foil faced", "-", "Rockwool / Equiv.", "m2", 28.0),
        ("Sub-Frame & Accessories", "Breather / vapour membrane", "Per m2", "-", "Imported", "m2", 12.0),
        ("Other Panels", "HPL exterior cladding panel", "8 mm", "Wood / colour", "Trespa / Fundermax / Equiv.", "m2", 320.0),
        ("Other Panels", "Fibre cement board", "8-12 mm", "-", "Local / Imported", "m2", 85.0),
        ("Other Panels", "Metal mesh / expanded metal facade", "Aluminium / SS", "-", "Imported", "m2", 420.0),
    ])

cat("10 Mild Steel", "MS", "Mild Steel (MS) - Sections, Plates & Sheets",
    "SHS, RHS, angles, channels, beams, flats, rounds, pipes, plates, GI sheets, grating", "D9D9D9", 0.07, [
        ("SHS (Square Hollow)", "MS square hollow section", "25 x 25 x 1.5 mm (1.06 kg/m)", "Black", "Local", "lm", ms(1.06)),
        ("SHS (Square Hollow)", "MS square hollow section", "40 x 40 x 2.0 mm (2.31 kg/m)", "Black", "Local", "lm", ms(2.31)),
        ("SHS (Square Hollow)", "MS square hollow section", "50 x 50 x 3.0 mm (4.25 kg/m)", "Black", "Local", "lm", ms(4.25)),
        ("SHS (Square Hollow)", "MS square hollow section", "60 x 60 x 3.0 mm (5.19 kg/m)", "Black", "Local", "lm", ms(5.19)),
        ("SHS (Square Hollow)", "MS square hollow section", "75 x 75 x 3.0 mm (6.71 kg/m)", "Black", "Local", "lm", ms(6.71)),
        ("SHS (Square Hollow)", "MS square hollow section", "100 x 100 x 4.0 mm (11.70 kg/m)", "Black", "Local", "lm", ms(11.70)),
        ("SHS (Square Hollow)", "MS square hollow section", "100 x 100 x 5.0 mm (14.40 kg/m)", "Black", "Local", "lm", ms(14.40)),
        ("SHS (Square Hollow)", "MS square hollow section", "150 x 150 x 6.0 mm (26.40 kg/m)", "Black", "Local", "lm", ms(26.40)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "50 x 25 x 2.0 mm (2.15 kg/m)", "Black", "Local", "lm", ms(2.15)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "80 x 40 x 3.0 mm (5.19 kg/m)", "Black", "Local", "lm", ms(5.19)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "100 x 50 x 3.0 mm (6.71 kg/m)", "Black", "Local", "lm", ms(6.71)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "100 x 50 x 4.0 mm (8.59 kg/m)", "Black", "Local", "lm", ms(8.59)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "150 x 100 x 5.0 mm (18.70 kg/m)", "Black", "Local", "lm", ms(18.70)),
        ("RHS (Rectangular Hollow)", "MS rectangular hollow section", "200 x 100 x 6.0 mm (27.30 kg/m)", "Black", "Local", "lm", ms(27.30)),
        ("Angles", "MS equal angle", "40 x 40 x 4 mm (2.42 kg/m)", "Black", "Local", "lm", ms(2.42)),
        ("Angles", "MS equal angle", "50 x 50 x 5 mm (3.77 kg/m)", "Black", "Local", "lm", ms(3.77)),
        ("Angles", "MS equal angle", "65 x 65 x 6 mm (5.91 kg/m)", "Black", "Local", "lm", ms(5.91)),
        ("Angles", "MS equal angle", "75 x 75 x 6 mm (6.85 kg/m)", "Black", "Local", "lm", ms(6.85)),
        ("Angles", "MS equal angle", "100 x 100 x 8 mm (12.20 kg/m)", "Black", "Local", "lm", ms(12.20)),
        ("Channels & Beams", "MS parallel flange channel (PFC)", "100 x 50 mm (10.20 kg/m)", "Black", "Local / Imported", "lm", ms(10.20)),
        ("Channels & Beams", "MS parallel flange channel (PFC)", "150 x 75 mm (17.90 kg/m)", "Black", "Local / Imported", "lm", ms(17.90)),
        ("Channels & Beams", "MS parallel flange channel (PFC)", "200 x 75 mm (23.40 kg/m)", "Black", "Local / Imported", "lm", ms(23.40)),
        ("Channels & Beams", "MS I-beam IPE 100", "8.10 kg/m", "Black", "Imported", "lm", ms(8.10)),
        ("Channels & Beams", "MS I-beam IPE 160", "15.80 kg/m", "Black", "Imported", "lm", ms(15.80)),
        ("Channels & Beams", "MS I-beam IPE 200", "22.40 kg/m", "Black", "Imported", "lm", ms(22.40)),
        ("Channels & Beams", "MS universal beam UB 203x133x25", "25.10 kg/m", "Black", "Imported", "lm", ms(25.10)),
        ("Channels & Beams", "MS universal beam UB 254x146x31", "31.10 kg/m", "Black", "Imported", "lm", ms(31.10)),
        ("Channels & Beams", "MS H-beam HEA 200", "42.30 kg/m", "Black", "Imported", "lm", ms(42.30)),
        ("Channels & Beams", "GI purlin C / Z section", "200 x 2.0 mm (5.60 kg/m)", "Pre-galvanized", "Local", "lm", f"=ROUND(5.6*GI_KG,2)"),
        ("Flats & Rounds", "MS flat bar", "25 x 3 mm (0.59 kg/m)", "Black", "Local", "lm", ms(0.59)),
        ("Flats & Rounds", "MS flat bar", "40 x 5 mm (1.57 kg/m)", "Black", "Local", "lm", ms(1.57)),
        ("Flats & Rounds", "MS flat bar", "50 x 6 mm (2.36 kg/m)", "Black", "Local", "lm", ms(2.36)),
        ("Flats & Rounds", "MS flat bar", "75 x 8 mm (4.71 kg/m)", "Black", "Local", "lm", ms(4.71)),
        ("Flats & Rounds", "MS flat bar", "100 x 10 mm (7.85 kg/m)", "Black", "Local", "lm", ms(7.85)),
        ("Flats & Rounds", "MS round bar", "Dia 12 mm (0.89 kg/m)", "Black", "Local", "lm", ms(0.89)),
        ("Flats & Rounds", "MS round bar", "Dia 16 mm (1.58 kg/m)", "Black", "Local", "lm", ms(1.58)),
        ("Flats & Rounds", "MS round bar", "Dia 20 mm (2.47 kg/m)", "Black", "Local", "lm", ms(2.47)),
        ("Flats & Rounds", "MS square bar", "12 x 12 mm (1.13 kg/m)", "Black", "Local", "lm", ms(1.13)),
        ("Pipes (CHS)", "MS circular hollow section", "48.3 x 3.2 mm (3.56 kg/m)", "Black", "Local", "lm", ms(3.56)),
        ("Pipes (CHS)", "MS circular hollow section", "60.3 x 3.6 mm (5.03 kg/m)", "Black", "Local", "lm", ms(5.03)),
        ("Pipes (CHS)", "MS circular hollow section", "88.9 x 4.0 mm (8.38 kg/m)", "Black", "Local", "lm", ms(8.38)),
        ("Pipes (CHS)", "MS circular hollow section", "114.3 x 4.5 mm (12.20 kg/m)", "Black", "Local", "lm", ms(12.20)),
        ("Plates", "MS plate", "3 mm (23.55 kg/m2)", "Black", "Local", "m2", msp(23.55)),
        ("Plates", "MS plate", "5 mm (39.25 kg/m2)", "Black", "Local", "m2", msp(39.25)),
        ("Plates", "MS plate", "6 mm (47.10 kg/m2)", "Black", "Local", "m2", msp(47.10)),
        ("Plates", "MS plate", "8 mm (62.80 kg/m2)", "Black", "Local", "m2", msp(62.80)),
        ("Plates", "MS plate", "10 mm (78.50 kg/m2)", "Black", "Local", "m2", msp(78.50)),
        ("Plates", "MS plate", "12 mm (94.20 kg/m2)", "Black", "Local", "m2", msp(94.20)),
        ("Plates", "MS plate", "16 mm (125.60 kg/m2)", "Black", "Local", "m2", msp(125.60)),
        ("Plates", "MS plate", "20 mm (157.00 kg/m2)", "Black", "Local", "m2", msp(157.00)),
        ("Plates", "MS chequered plate", "4.5 mm (39.50 kg/m2)", "Black", "Local", "m2", msp(39.50)),
        ("Sheets", "MS sheet (CR)", "1.5 mm (11.78 kg/m2)", "Black", "Local", "m2", msp(11.78)),
        ("Sheets", "MS sheet (CR)", "2.0 mm (15.70 kg/m2)", "Black", "Local", "m2", msp(15.70)),
        ("Sheets", "GI sheet", "0.8 mm (6.28 kg/m2)", "Galvanized", "Local", "m2", f"=ROUND(6.28*GI_KG,2)"),
        ("Sheets", "GI sheet", "1.2 mm (9.42 kg/m2)", "Galvanized", "Local", "m2", f"=ROUND(9.42*GI_KG,2)"),
        ("Sheets", "GI sheet", "1.5 mm (11.78 kg/m2)", "Galvanized", "Local", "m2", f"=ROUND(11.78*GI_KG,2)"),
        ("Sheets", "GI sheet", "2.0 mm (15.70 kg/m2)", "Galvanized", "Local", "m2", f"=ROUND(15.70*GI_KG,2)"),
        ("Sheets", "MS perforated sheet", "2.0 mm", "Black", "Local", "m2", 95.0),
        ("Sheets", "MS expanded metal mesh", "Standard", "Black / GI", "Local", "m2", 65.0),
        ("Grating & Misc.", "MS bar grating", "25 x 3 mm, 30 x 100 pitch", "HDG", "Local", "m2", 180.0),
        ("Grating & Misc.", "MS base plate with anchor holes", "200 x 200 x 10 mm", "Black / HDG", "Local fabrication", "nos", 25.0),
        ("Grating & Misc.", "MS embed plate with studs", "150 x 150 x 10 mm", "HDG", "Local fabrication", "nos", 40.0),
        ("Grating & Misc.", "MS pipe railing / balustrade (material only)", "Per lm", "Black", "Local fabrication", "lm", 85.0),
    ])

cat("11 Stainless Steel", "SS", "Stainless Steel (SS) - Tubes, Sheets & Drains",
    "SS 304 / 316 round, square & rectangular tubes, sheets, flats, rods, linear drains", "BFBFBF", 0.05, [
        ("Round Tube", "SS round tube", "25.4 x 1.2 mm (0.73 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(0.73)),
        ("Round Tube", "SS round tube", "38.1 x 1.5 mm (1.37 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.37)),
        ("Round Tube", "SS round tube", "42.4 x 1.5 mm (1.53 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.53)),
        ("Round Tube", "SS round tube", "50.8 x 1.5 mm (1.85 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.85)),
        ("Round Tube", "SS round tube", "50.8 x 2.0 mm (2.44 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(2.44)),
        ("Round Tube", "SS round tube", "63.5 x 2.0 mm (3.07 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(3.07)),
        ("Round Tube", "SS round tube", "76.2 x 2.0 mm (3.71 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(3.71)),
        ("Round Tube", "SS round tube", "42.4 x 2.0 mm (2.02 kg/m)", "SS 316 hairline", "Local / Imported", "lm", ss(2.02, "SS316")),
        ("Round Tube", "SS round tube", "50.8 x 2.0 mm (2.44 kg/m)", "SS 316 hairline", "Local / Imported", "lm", ss(2.44, "SS316")),
        ("Square / Rect. Tube", "SS square tube", "25 x 25 x 1.2 mm (0.92 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(0.92)),
        ("Square / Rect. Tube", "SS square tube", "40 x 40 x 1.5 mm (1.83 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.83)),
        ("Square / Rect. Tube", "SS square tube", "50 x 50 x 1.5 mm (2.31 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(2.31)),
        ("Square / Rect. Tube", "SS square tube", "50 x 50 x 2.0 mm (3.05 kg/m)", "SS 316 hairline", "Local / Imported", "lm", ss(3.05, "SS316")),
        ("Square / Rect. Tube", "SS rectangular tube", "40 x 20 x 1.2 mm (1.10 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.10)),
        ("Square / Rect. Tube", "SS rectangular tube", "50 x 25 x 1.5 mm (1.73 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(1.73)),
        ("Square / Rect. Tube", "SS rectangular tube", "60 x 40 x 1.5 mm (2.31 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(2.31)),
        ("Square / Rect. Tube", "SS rectangular tube", "100 x 50 x 2.0 mm (4.59 kg/m)", "SS 304 hairline", "Local / Imported", "lm", ss(4.59)),
        ("Flats, Rods & Angles", "SS flat bar", "40 x 5 mm (1.59 kg/m)", "SS 304", "Local / Imported", "lm", ss(1.59)),
        ("Flats, Rods & Angles", "SS flat bar", "50 x 10 mm (3.97 kg/m)", "SS 304", "Local / Imported", "lm", ss(3.97)),
        ("Flats, Rods & Angles", "SS round rod", "Dia 10 mm (0.62 kg/m)", "SS 304", "Local / Imported", "lm", ss(0.62)),
        ("Flats, Rods & Angles", "SS round rod", "Dia 12 mm (0.89 kg/m)", "SS 304", "Local / Imported", "lm", ss(0.89)),
        ("Flats, Rods & Angles", "SS equal angle", "40 x 40 x 4 mm (2.45 kg/m)", "SS 304", "Local / Imported", "lm", ss(2.45)),
        ("Sheets", "SS sheet", "1.0 mm (7.93 kg/m2)", "SS 304 hairline (#4)", "Local / Imported", "m2", ss(7.93)),
        ("Sheets", "SS sheet", "1.2 mm (9.52 kg/m2)", "SS 304 hairline (#4)", "Local / Imported", "m2", ss(9.52)),
        ("Sheets", "SS sheet", "1.5 mm (11.90 kg/m2)", "SS 304 hairline (#4)", "Local / Imported", "m2", ss(11.90)),
        ("Sheets", "SS sheet", "2.0 mm (15.86 kg/m2)", "SS 304 hairline (#4)", "Local / Imported", "m2", ss(15.86)),
        ("Sheets", "SS sheet", "3.0 mm (23.79 kg/m2)", "SS 304 hairline (#4)", "Local / Imported", "m2", ss(23.79)),
        ("Sheets", "SS sheet", "1.5 mm (11.90 kg/m2)", "SS 304 mirror (#8)", "Local / Imported", "m2", "=ROUND(11.9*(SS304+SS_MIRROR),2)"),
        ("Sheets", "SS sheet", "1.5 mm (11.90 kg/m2)", "SS 316 hairline (#4)", "Local / Imported", "m2", ss(11.90, "SS316")),
        ("Sheets", "SS sheet", "2.0 mm (15.86 kg/m2)", "SS 316 hairline (#4)", "Local / Imported", "m2", ss(15.86, "SS316")),
        ("Sheets", "SS PVD coloured sheet", "1.2 mm", "Gold / Rose gold / Black", "Imported", "m2", 420.0),
        ("Sheets", "SS chequered / anti-slip plate", "3.0 mm", "SS 304", "Imported", "m2", 520.0),
        ("Linear Drains", "Linear shower drain - tile insert", "600 mm, SS 304", "Satin", "Local / Imported", "nos", 250.0),
        ("Linear Drains", "Linear shower drain - tile insert", "800 mm, SS 304", "Satin", "Local / Imported", "nos", 300.0),
        ("Linear Drains", "Linear shower drain - tile insert", "1000 mm, SS 304", "Satin", "Local / Imported", "nos", 380.0),
        ("Linear Drains", "Linear drain - slotted / grated cover", "800 mm, SS 304", "Satin", "Local / Imported", "nos", 280.0),
        ("Linear Drains", "Linear drain - custom fabricated", "SS 316, 1.5 mm, per lm", "Satin", "Local fabrication", "lm", 450.0),
        ("Linear Drains", "Heavy-duty trench drain grating", "SS 304, 150 mm wide", "Satin", "Local fabrication", "lm", 650.0),
        ("Misc. SS Items", "SS corner guard / edge trim", "1.0 mm, 50 x 50", "SS 304 hairline", "Local", "lm", 45.0),
        ("Misc. SS Items", "SS skirting / kick plate", "1.2 mm, 150 mm high", "SS 304 hairline", "Local", "lm", 60.0),
        ("Misc. SS Items", "SS tile trim / profile", "10-12 mm", "SS 304", "Local / Imported", "lm", 22.0),
        ("Misc. SS Items", "SS access / manhole cover (recessed)", "600 x 600 mm", "SS 304", "Local / Imported", "nos", 950.0),
    ])

cat("12 Auto Doors & Shutters", "AD", "Automatic Doors, Operators & Shutters",
    "Sliding / swing operators, sensors, access control, rolling shutters", "B4C6E7", 0.0, [
        ("Automatic Sliding Doors", "Automatic sliding door operator", "Single / double leaf, up to 2 x 120 kg", "Silver", "dormakaba ES 200 / Equiv.", "set", 10500.0),
        ("Automatic Sliding Doors", "Automatic sliding door operator", "Single / double leaf, up to 2 x 100 kg", "Silver", "GEZE Slimdrive / Equiv.", "set", 11000.0),
        ("Automatic Sliding Doors", "Automatic sliding door operator", "Up to 2 x 100 kg", "Silver", "Local / China", "set", 5500.0),
        ("Automatic Sliding Doors", "Telescopic sliding operator", "4 leaf", "Silver", "dormakaba / GEZE / Equiv.", "set", 16500.0),
        ("Automatic Sliding Doors", "Hermetic sliding door operator (hospital)", "Single leaf", "-", "Imported", "set", 18000.0),
        ("Automatic Swing Doors", "Automatic swing door operator", "Single leaf, up to 250 kg", "Silver", "dormakaba ED 100 / Equiv.", "set", 8500.0),
        ("Automatic Swing Doors", "Automatic swing door operator", "Single leaf", "Silver", "Local / China", "set", 4200.0),
        ("Sensors & Accessories", "Motion / presence combined sensor", "-", "-", "BEA / Equiv.", "nos", 650.0),
        ("Sensors & Accessories", "Safety beam / photocell", "-", "-", "BEA / Equiv.", "nos", 250.0),
        ("Sensors & Accessories", "Programme switch / key switch", "-", "-", "OEM", "nos", 350.0),
        ("Sensors & Accessories", "Battery back-up unit", "-", "-", "OEM", "nos", 750.0),
        ("Sensors & Accessories", "Push / touchless exit button", "-", "SS", "Local", "nos", 85.0),
        ("Sensors & Accessories", "Access control (card / keypad) reader", "Standalone", "-", "Local / Imported", "nos", 450.0),
        ("Rolling Shutters", "Aluminium rolling shutter - manual", "Per m2", "PC", "Local", "m2", 280.0),
        ("Rolling Shutters", "Aluminium rolling shutter - motorized", "Per m2 (with motor)", "PC", "Local", "m2", 420.0),
        ("Rolling Shutters", "GI rolling shutter - motorized", "Per m2", "Galvanized / PC", "Local", "m2", 320.0),
        ("Rolling Shutters", "Fire-rated rolling shutter", "Per m2, 2 hr", "Galvanized", "Imported", "m2", 950.0),
        ("Rolling Shutters", "Tubular motor for shutter", "Up to 50 Nm", "-", "Somfy / Equiv.", "nos", 950.0),
    ])

cat("13 Coating & Finishing", "CF", "Coating, Finishing & Treatment Services",
    "Powder coating, anodizing, PVDF, galvanizing, painting, polishing (job-work rates)", "E2EFDA", 0.0, [
        ("Aluminium Finishing", "Powder coating (job work)", "Standard RAL, 60-80 micron", "Matt / gloss", "Local applicator", "kg", 3.0),
        ("Aluminium Finishing", "Powder coating (job work)", "Standard RAL", "Matt / gloss", "Local applicator", "m2", 30.0),
        ("Aluminium Finishing", "Powder coating - special / metallic / texture", "Per m2", "Special", "Local applicator", "m2", 45.0),
        ("Aluminium Finishing", "Anodizing (job work)", "15-20 micron", "Natural / bronze / black", "Local applicator", "kg", 5.0),
        ("Aluminium Finishing", "PVDF coating (job work)", "2-coat, 35-40 micron", "Colour as approved", "Local applicator", "m2", 70.0),
        ("Aluminium Finishing", "PVDF coating (job work)", "3-coat metallic", "Metallic", "Local applicator", "m2", 90.0),
        ("Aluminium Finishing", "Wood-grain sublimation", "Over powder coat", "Wood effect", "Local applicator", "m2", 55.0),
        ("Steel Finishing", "Hot-dip galvanizing", "Per kg of steel", "HDG", "Local galvanizer", "kg", 1.6),
        ("Steel Finishing", "Sand blasting (SA 2.5)", "Per m2", "-", "Local", "m2", 20.0),
        ("Steel Finishing", "Primer + enamel paint (2 coats)", "Per m2", "Colour as req.", "Jotun / National / Equiv.", "m2", 25.0),
        ("Steel Finishing", "Epoxy / PU paint system (3 coat)", "Per m2", "Colour as req.", "Jotun / Hempel / Equiv.", "m2", 45.0),
        ("Steel Finishing", "Intumescent fire-protective paint", "60 min, per m2", "-", "Jotun / Equiv.", "m2", 120.0),
        ("Steel Finishing", "Powder coating on MS / GI", "Per m2", "RAL", "Local applicator", "m2", 35.0),
        ("Stainless Steel Finishing", "Hairline re-polishing / weld polishing", "Per lm of weld", "#4", "In-house / Local", "lm", 15.0),
        ("Stainless Steel Finishing", "Mirror polishing", "Per m2", "#8", "Local", "m2", 120.0),
        ("Stainless Steel Finishing", "PVD colour coating", "Per m2", "Gold / Black / Rose gold", "Local / Imported", "m2", 260.0),
        ("Stainless Steel Finishing", "Passivation / pickling", "Per m2", "-", "Local", "m2", 25.0),
    ])

cat("14 Labour & Equipment", "LB", "Labour, Equipment, Transport & Misc. Costs",
    "Manpower day rates, access equipment, cranes, transport, scaffolding, testing", "FCE4D6", 0.0, [
        ("Manpower (cost per day)", "Project manager", "8 hr day, incl. visa & benefits", "-", "In-house", "day", 650.0),
        ("Manpower (cost per day)", "Site engineer", "8 hr day, incl. visa & benefits", "-", "In-house", "day", 450.0),
        ("Manpower (cost per day)", "Foreman / supervisor", "8 hr day, incl. visa & benefits", "-", "In-house", "day", 300.0),
        ("Manpower (cost per day)", "Aluminium fabricator", "8 hr day", "-", "In-house", "day", 220.0),
        ("Manpower (cost per day)", "Aluminium installer", "8 hr day", "-", "In-house", "day", 220.0),
        ("Manpower (cost per day)", "Glazier / glass fixer", "8 hr day", "-", "In-house", "day", 230.0),
        ("Manpower (cost per day)", "Steel fabricator / welder", "8 hr day", "-", "In-house", "day", 250.0),
        ("Manpower (cost per day)", "SS (TIG) welder / polisher", "8 hr day", "-", "In-house", "day", 280.0),
        ("Manpower (cost per day)", "Painter", "8 hr day", "-", "In-house", "day", 180.0),
        ("Manpower (cost per day)", "Helper", "8 hr day", "-", "In-house", "day", 140.0),
        ("Manpower (cost per day)", "Driver", "8 hr day", "-", "In-house", "day", 180.0),
        ("Manpower (cost per day)", "Overtime premium", "Per man-hour", "-", "In-house", "hr", 25.0),
        ("Manpower (cost per day)", "Hired labour (sub-contract)", "8 hr day", "-", "Labour supply co.", "day", 160.0),
        ("Access Equipment", "Boom lift 45 ft", "Daily hire with operator", "-", "Rental co.", "day", 650.0),
        ("Access Equipment", "Boom lift 60-80 ft", "Daily hire with operator", "-", "Rental co.", "day", 1100.0),
        ("Access Equipment", "Scissor lift 19-26 ft", "Daily hire", "-", "Rental co.", "day", 350.0),
        ("Access Equipment", "Mobile scaffolding tower", "Monthly hire", "-", "Rental co.", "month", 900.0),
        ("Access Equipment", "Scaffolding (erected)", "Per m2 per month", "-", "Scaffolding contractor", "m2", 12.0),
        ("Access Equipment", "Gondola / BMU hire", "Monthly", "-", "Rental co.", "month", 6500.0),
        ("Access Equipment", "Rope access team (2 men)", "Per day", "-", "Rope access co.", "day", 1500.0),
        ("Lifting & Handling", "Mobile crane 25 ton", "8 hr with operator", "-", "Rental co.", "day", 1800.0),
        ("Lifting & Handling", "Mobile crane 50 ton", "8 hr with operator", "-", "Rental co.", "day", 3200.0),
        ("Lifting & Handling", "Glass vacuum lifter (manual)", "Daily hire", "-", "Rental co.", "day", 150.0),
        ("Lifting & Handling", "Glass vacuum lifter (mini crane / robot)", "Daily hire with operator", "-", "Rental co.", "day", 1600.0),
        ("Lifting & Handling", "Chain block / manual hoist", "Daily hire", "-", "Rental co.", "day", 60.0),
        ("Transport", "Pickup trip (within city)", "Per trip", "-", "In-house", "trip", 200.0),
        ("Transport", "3-ton truck trip (within city)", "Per trip", "-", "In-house / Hired", "trip", 400.0),
        ("Transport", "Flatbed trailer trip (inter-emirate)", "Per trip", "-", "Hired", "trip", 1500.0),
        ("Transport", "Glass A-frame truck trip", "Per trip", "-", "Hired", "trip", 800.0),
        ("Testing & Approvals", "Site water test (AAMA 501.2)", "Per test", "-", "Third-party lab", "nos", 1500.0),
        ("Testing & Approvals", "Anchor pull-out test", "Per anchor", "-", "Third-party lab", "nos", 250.0),
        ("Testing & Approvals", "Structural calculations & approval", "Per project", "-", "Consultant", "ls", 6000.0),
        ("Testing & Approvals", "Civil defence approval / certification", "Per project", "-", "Authority / Consultant", "ls", 3500.0),
        ("Testing & Approvals", "Shop drawings preparation", "Per sheet", "-", "In-house / Outsourced", "sheet", 150.0),
        ("Testing & Approvals", "Material samples & mock-up", "Per project", "-", "In-house", "ls", 2500.0),
        ("Site Overheads", "Site storage / container", "Monthly", "-", "Rental co.", "month", 800.0),
        ("Site Overheads", "Protection & cleaning on completion", "Per m2", "-", "In-house", "m2", 5.0),
        ("Site Overheads", "Safety (PPE, signage, barricades)", "Per project", "-", "In-house", "ls", 1500.0),
        ("Site Overheads", "Insurance (CAR / third party)", "% of contract value", "-", "Insurer", "%", 0.5),
        ("Site Overheads", "Performance / advance bank guarantee charges", "% p.a. of guarantee", "-", "Bank", "%", 1.5),
    ])

cat("15 System Rates S&I", "SI", "Finished Systems - Budget Supply & Install Rates",
    "All-in budget rates for quick estimates (material + fabrication + installation)", "D0CECE", 0.0, [
        ("Curtain Wall & Facade", "Stick curtain wall - clear DGU", "Gulf Extrusions / equiv., 6+12+6 DGU", "PC / PVDF", "Local system", "m2", 1050.0),
        ("Curtain Wall & Facade", "Stick curtain wall - Low-E DGU", "Gulf Extrusions / equiv., 6+12+6 Low-E", "PVDF", "Local system", "m2", 1250.0),
        ("Curtain Wall & Facade", "Unitized curtain wall", "Low-E DGU, PVDF", "PVDF", "Local system", "m2", 1650.0),
        ("Curtain Wall & Facade", "Structural glazing (SSG) curtain wall", "Low-E DGU", "PVDF", "Local system", "m2", 1400.0),
        ("Curtain Wall & Facade", "Spider (point-fixed) glazing", "SS 316 spiders, 12.76-21.52 TL", "-", "Local / Imported fittings", "m2", 2200.0),
        ("Curtain Wall & Facade", "Glass fin wall", "Laminated glass fins", "-", "Local", "m2", 2600.0),
        ("Curtain Wall & Facade", "ACP cladding - supply & install", "4 mm FR, with sub-frame", "PVDF", "Alubond / Equiv.", "m2", 220.0),
        ("Curtain Wall & Facade", "Solid aluminium panel cladding - S&I", "3 mm, with sub-frame", "PVDF", "Local", "m2", 380.0),
        ("Windows", "Sliding window - non-thermal", "6 mm tempered / 6.38 lam", "PC", "Gulf Extrusions / Equiv.", "m2", 480.0),
        ("Windows", "Sliding window - thermal break", "DGU 6+12+6", "PC", "Gulf Extrusions / Equiv.", "m2", 750.0),
        ("Windows", "Casement / top-hung window - non-thermal", "6 mm tempered", "PC", "Gulf Extrusions / Equiv.", "m2", 550.0),
        ("Windows", "Casement / top-hung window - thermal break", "DGU 6+12+6", "PC", "Gulf Extrusions / Equiv.", "m2", 820.0),
        ("Windows", "Fixed window - thermal break", "DGU 6+12+6", "PC", "Gulf Extrusions / Equiv.", "m2", 620.0),
        ("Windows", "Tilt & turn window", "DGU, premium system", "PC", "Schuco / Technal / Reynaers", "m2", 1650.0),
        ("Windows", "Slim-line minimal window / door", "DGU, premium system", "PC", "Imported", "m2", 2800.0),
        ("Doors", "Aluminium hinged door - single", "Glazed, with closer & lock", "PC", "Gulf Extrusions / Equiv.", "m2", 950.0),
        ("Doors", "Aluminium hinged door - thermal break", "DGU, multipoint lock", "PC", "Gulf Extrusions / Equiv.", "m2", 1250.0),
        ("Doors", "Aluminium sliding door - thermal break", "DGU 6+12+6", "PC", "Gulf Extrusions / Equiv.", "m2", 820.0),
        ("Doors", "Lift & slide door", "DGU, premium system", "PC", "Schuco / Technal / Equiv.", "m2", 1900.0),
        ("Doors", "Folding / bi-fold door", "DGU", "PC", "Local / Imported system", "m2", 1500.0),
        ("Doors", "Frameless glass door - 12 mm with patch fittings", "Floor spring, H-handle", "SS", "Local fittings", "nos", 3200.0),
        ("Doors", "Automatic sliding glass door (complete)", "Double leaf, 2.0 x 2.4 m", "PC", "dormakaba / GEZE", "nos", 16500.0),
        ("Doors", "Fire-rated glazed aluminium door", "60 min", "PC", "Imported system", "m2", 3800.0),
        ("Partitions", "Frameless glass partition", "12 mm tempered, U-channel", "Anodized", "Local", "m2", 420.0),
        ("Partitions", "Framed aluminium glass partition", "Single glazed 10 mm", "Anodized / PC", "Local", "m2", 550.0),
        ("Partitions", "Double-glazed acoustic partition", "2 x 10.76 lam", "Anodized", "Local / Imported", "m2", 1100.0),
        ("Partitions", "Back-painted glass wall cladding", "6 mm tempered", "RAL", "Local", "m2", 280.0),
        ("Balustrades & Handrails", "Glass balustrade with SS posts & clamps", "12 mm tempered, SS 304, 1.1 m H", "SS satin", "Local", "lm", 950.0),
        ("Balustrades & Handrails", "Frameless glass balustrade - base shoe", "17.52 TL, aluminium shoe, 1.1 m H", "Anodized / SS clad", "Local", "lm", 1450.0),
        ("Balustrades & Handrails", "Glass balustrade - spigot mounted", "15 mm tempered / 17.52 TL", "SS 316", "Local / Imported", "lm", 1250.0),
        ("Balustrades & Handrails", "SS handrail - wall mounted", "Dia 50.8 mm, SS 304", "Hairline", "Local", "lm", 280.0),
        ("Balustrades & Handrails", "SS balustrade with rod / wire infill", "SS 316, 1.1 m H", "Hairline", "Local", "lm", 850.0),
        ("Balustrades & Handrails", "MS balustrade / railing", "Painted, 1.1 m H", "Paint / PC", "Local", "lm", 450.0),
        ("Balustrades & Handrails", "Aluminium balustrade with vertical pickets", "1.1 m H", "PC", "Local", "lm", 550.0),
        ("Shower & Bathroom", "Frameless shower enclosure - fixed panel", "10 mm tempered", "Chrome / satin fittings", "Local", "m2", 550.0),
        ("Shower & Bathroom", "Frameless shower enclosure - hinged door + fixed", "10 mm tempered, up to 1.2 x 2.0 m", "Chrome / satin fittings", "Local", "set", 2200.0),
        ("Shower & Bathroom", "Sliding shower enclosure", "8-10 mm tempered", "SS fittings", "Local", "set", 2600.0),
        ("Shower & Bathroom", "Bathtub screen", "8 mm tempered, hinged", "Chrome", "Local", "nos", 950.0),
        ("Shower & Bathroom", "Mirror supply & fix", "6 mm polished edge", "Clear silver", "Local", "m2", 160.0),
        ("Shower & Bathroom", "Linear shower drain supply & fix", "800 mm, SS 304", "Satin", "Local / Imported", "nos", 450.0),
        ("Pergola, Canopy & Skylight", "Aluminium pergola (fixed slats)", "Per m2 plan area", "PC / wood-grain", "Local", "m2", 650.0),
        ("Pergola, Canopy & Skylight", "Bioclimatic motorized louvered pergola", "Per m2 plan area", "PC", "Local / Imported", "m2", 1500.0),
        ("Pergola, Canopy & Skylight", "MS pergola", "Painted", "Paint", "Local", "m2", 450.0),
        ("Pergola, Canopy & Skylight", "Glass canopy with SS tie-rods", "21.52 TL", "SS 316", "Local / Imported", "m2", 1650.0),
        ("Pergola, Canopy & Skylight", "Polycarbonate canopy on MS / aluminium frame", "10 mm multiwall", "PC", "Local", "m2", 380.0),
        ("Pergola, Canopy & Skylight", "Skylight - aluminium frame, Low-E DGU", "Laminated inner", "PC", "Local", "m2", 1350.0),
        ("Pergola, Canopy & Skylight", "Walk-on glass floor / skylight", "Triple laminated, anti-slip", "SS / MS frame", "Local", "m2", 2800.0),
        ("Louvers & Screens", "Aluminium fixed louvers", "Z / aerofoil blades", "PC", "Local", "m2", 420.0),
        ("Louvers & Screens", "Aluminium operable louvers (motorized)", "Aerofoil blades", "PC", "Imported", "m2", 1200.0),
        ("Louvers & Screens", "Sand-trap / weather louvers", "With bird / insect mesh", "PC", "Local", "m2", 550.0),
        ("Louvers & Screens", "Decorative laser-cut screen", "3 mm aluminium", "PVDF", "Local", "m2", 650.0),
        ("MS & SS Works", "MS structural steel - fabricated & erected", "Painted, per kg", "Paint", "Local", "kg", 9.5),
        ("MS & SS Works", "MS structural steel - fabricated & erected", "Hot-dip galvanized, per kg", "HDG", "Local", "kg", 11.5),
        ("MS & SS Works", "MS staircase (stringers + chequered treads)", "Per riser", "Paint", "Local", "nos", 650.0),
        ("MS & SS Works", "SS fabrication - general", "SS 304, per kg", "Hairline", "Local", "kg", 55.0),
        ("MS & SS Works", "SS cladding (column / wall)", "1.2 mm SS 304 on sub-frame", "Hairline", "Local", "m2", 650.0),
    ])

# ---------------------------------------------------------------- suppliers
SUPPLIERS = [
    ("Aluminium Extrusions", "Gulf Extrusions Co. LLC", "Profiles, systems, coating"),
    ("Aluminium Extrusions", "White Aluminium Extrusion (WAE)", "Profiles, systems"),
    ("Aluminium Extrusions", "Emirates Extrusion Factory (EEF)", "Profiles"),
    ("Aluminium Systems", "Schuco / Technal / Reynaers (via fabricators)", "Premium window, door & facade systems"),
    ("Float Glass", "Emirates Float Glass (EFG)", "Float, tinted, reflective, coated"),
    ("Float / Coated Glass", "Guardian Glass (RAK)", "Float, SunGuard Low-E, mirror"),
    ("Glass Processing", "Emirates Glass LLC", "Tempered, laminated, DGU, Low-E"),
    ("Glass Processing", "Dubai Glass Industries / Equiv.", "Tempered, laminated, DGU"),
    ("Door Hardware", "dormakaba Middle East", "Floor springs, closers, patch fittings, automatic doors"),
    ("Door Hardware", "GEZE Middle East", "Closers, automatic doors"),
    ("Hardware", "Häfele Middle East", "Architectural & shower hardware"),
    ("Window Hardware", "Giesse / Savio / Roto / Siegenia (distributors)", "Window & door hardware"),
    ("Sealants", "Dow (DOWSIL) distributors", "Structural & weather silicones"),
    ("Sealants", "Sika Gulf", "Sealants, adhesives, PU foam"),
    ("Anchors", "Hilti Emirates / Fischer", "Anchors, chemical anchors, firestop"),
    ("ACP / Cladding", "Alubond U.S.A. (Mulk Holdings)", "ACP FR & A2"),
    ("ACP / Cladding", "Alucobond / 3A (distributors)", "ACP"),
    ("Steel", "Local steel traders (Al Quoz / Sharjah Ind.)", "MS sections, plates, sheets"),
    ("Stainless Steel", "Local SS traders (Sharjah / Ajman)", "SS tubes, sheets, fittings"),
    ("Coating", "Local Qualicoat applicators", "Powder coating, anodizing, PVDF"),
    ("Galvanizing", "Local HDG plants", "Hot-dip galvanizing"),
    ("Equipment Rental", "Local rental companies", "Boom lifts, cranes, scaffolding"),
]

# ---------------------------------------------------------------- helpers
thin = Side(style="thin", color="BFC5D2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
F = lambda **k: Font(name=FONT, **k)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
NUM = '#,##0.00;(#,##0.00);"-"'
STATUS_LIST = '"Indicative,Quoted,Verified,Obsolete"'

wb = Workbook()
ws0 = wb.active
ws0.title = "Index"

# Settings sheet first (names needed by formulas)
st = wb.create_sheet("Settings")
st.sheet_properties.tabColor = GOLD
st["B2"] = "SETTINGS & BASE RATES"
st["B2"].font = F(bold=True, size=16, color=NAVY)
st["B3"] = "Change the yellow cells - every linked rate in the workbook updates automatically."
st["B3"].font = F(italic=True, size=9, color="595959")
for c, h in zip("BCDEF", ["Parameter", "Value", "Unit", "Name (used in formulas)", "Note"]):
    st[f"{c}5"] = h
    st[f"{c}5"].font = F(bold=True, color="FFFFFF")
    st[f"{c}5"].fill = HDR_FILL
    st[f"{c}5"].alignment = CENTER
    st[f"{c}5"].border = BORDER
for i, (name, label, val, unit, note) in enumerate(SETTINGS, start=6):
    st[f"B{i}"] = label
    st[f"C{i}"] = val
    st[f"C{i}"].font = F(color="0000FF", bold=True)
    st[f"C{i}"].fill = INPUT_FILL
    st[f"C{i}"].number_format = "0.0%" if unit == "%" else "#,##0.00"
    st[f"D{i}"] = unit
    st[f"E{i}"] = name
    st[f"F{i}"] = note
    for c in "BCDEF":
        st[f"{c}{i}"].border = BORDER
        if c != "C":
            st[f"{c}{i}"].font = F(size=10)
        st[f"{c}{i}"].alignment = LEFT if c in "BF" else CENTER
    wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!$C${i}")
last = 6 + len(SETTINGS)
st[f"B{last + 1}"] = "Rate validity / last reviewed"
st[f"C{last + 1}"] = "09-Oct-2026"
st[f"C{last + 1}"].font = F(color="0000FF", bold=True)
st[f"C{last + 1}"].fill = INPUT_FILL
st[f"B{last + 2}"] = "Currency"
st[f"C{last + 2}"] = "AED"
for r in (last + 1, last + 2):
    for c in "BC":
        st[f"{c}{r}"].border = BORDER
        st[f"{c}{r}"].alignment = LEFT if c == "B" else CENTER
    st[f"B{r}"].font = F(size=10)
st[f"B{last + 4}"] = ("Note: all base rates are INDICATIVE UAE market budget rates (Oct-2026) for estimation "
                      "reference only. Confirm with current supplier quotations before final pricing.")
st[f"B{last + 4}"].font = F(italic=True, size=9, color="C00000")
st.merge_cells(f"B{last + 4}:F{last + 5}")
st[f"B{last + 4}"].alignment = LEFT
for c, w in zip("ABCDEF", [2, 40, 14, 10, 24, 44]):
    st.column_dimensions[c].width = w
st.sheet_view.showGridLines = False

# ---------------------------------------------------------------- category sheets
HEAD = ["Item Code", "Sub-Category", "Item Description", "Specification / Size", "Finish / Colour",
        "Brand / Origin", "Unit", "Basic Rate\n(AED)", "Wastage\n%", "Net Cost Rate\n(AED)",
        "Selling Rate\n(AED, +O&P)", "Selling Rate\nincl. VAT (AED)", "Rate Status", "Supplier",
        "Quote Ref.", "Rate Date", "Remarks"]
WIDTHS = [10, 22, 40, 32, 20, 26, 7, 12, 9, 13, 13, 14, 12, 20, 13, 12, 26]
FIRST = 5
DB_ROWS = []  # (sheet, row) for consolidated sheet

for ci, c in enumerate(CATS, start=1):
    ws = wb.create_sheet(c["sheet"])
    ws.sheet_properties.tabColor = c["colour"]
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:Q1")
    ws["A1"] = f"{ci:02d}.  {c['title'].upper()}"
    ws["A1"].font = F(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = HDR_FILL
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:Q2")
    ws["A2"] = f"Scope: {c['scope']}"
    ws["A2"].font = F(italic=True, size=9, color="404040")
    ws["A2"].fill = PatternFill("solid", fgColor=c["colour"])
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A3"] = '=HYPERLINK("#Index!A1","<< Back to Index")'
    ws["A3"].font = F(size=9, color="0563C1", underline="single")
    ws["D3"] = "Items:"
    ws["D3"].font = F(size=9, bold=True)
    ws["D3"].alignment = Alignment(horizontal="right")
    n_items = len(c["items"])
    end = FIRST + n_items + SPARE_ROWS - 1
    ws["E3"] = f"=COUNTA(C{FIRST}:C{end})"
    ws["E3"].font = F(size=9, bold=True)
    ws["E3"].alignment = Alignment(horizontal="left")
    ws["H3"] = "Blue = input   |   Black = formula   |   Yellow rows at bottom = spare for new items"
    ws["H3"].font = F(size=8, italic=True, color="595959")
    for j, h in enumerate(HEAD, start=1):
        cell = ws.cell(row=4, column=j, value=h)
        cell.font = F(bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[4].height = 32

    for k in range(n_items + SPARE_ROWS):
        r = FIRST + k
        spare = k >= n_items
        if not spare:
            it = c["items"][k]
            sub, desc, spec, fin, brand, unit, rate = it[:7]
            wst = it[7] if len(it) > 7 else c["wastage"]
            ws.cell(r, 1, f"{c['prefix']}-{k + 1:03d}")
            for j, v in zip(range(2, 8), (sub, desc, spec, fin, brand, unit)):
                ws.cell(r, j, v)
            ws.cell(r, 8, rate)
            ws.cell(r, 9, wst)
            ws.cell(r, 13, "Indicative")
        else:
            ws.cell(r, 1, f'=IF(C{r}="","","{c["prefix"]}-"&TEXT(ROW()-{FIRST - 1},"000"))')
        ws.cell(r, 10, f'=IF(H{r}="","",ROUND(H{r}*(1+I{r}),2))')
        ws.cell(r, 11, f'=IF(J{r}="","",ROUND(J{r}*(1+OHP),2))')
        ws.cell(r, 12, f'=IF(K{r}="","",ROUND(K{r}*(1+VAT),2))')
        for j in range(1, 18):
            cell = ws.cell(r, j)
            cell.border = BORDER
            cell.alignment = LEFT if j in (2, 3, 4, 5, 6, 14, 17) else CENTER
            blue = j in (2, 3, 4, 5, 6, 7, 9, 13, 14, 15, 16, 17) or (j == 8 and not str(cell.value or "").startswith("="))
            cell.font = F(size=10, color="0000FF" if blue else "000000", bold=(j == 3 and not spare))
            if j in (8, 10, 11, 12):
                cell.number_format = NUM
            if j == 9:
                cell.number_format = "0%"
            if j == 16:
                cell.number_format = "dd-mmm-yy"
            if spare:
                cell.fill = INPUT_FILL
        if not spare and str(ws.cell(r, 8).value).startswith("="):
            ws.cell(r, 8).font = F(size=10, color="008000")  # linked to Settings
        DB_ROWS.append((c["sheet"], c["title"], r))

    rng = f"A{FIRST}:Q{end}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND($C{FIRST}<>"",MOD(ROW(),2)=0,$M{FIRST}<>"Obsolete")'], fill=PatternFill("solid", fgColor=GREY)))
    for status, colour in (("Indicative", "FCE4D6"), ("Quoted", "DDEBF7"), ("Verified", "E2EFDA"), ("Obsolete", "D9D9D9")):
        ws.conditional_formatting.add(f"M{FIRST}:M{end}", FormulaRule(
            formula=[f'$M{FIRST}="{status}"'], fill=PatternFill("solid", fgColor=colour)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'$M{FIRST}="Obsolete"'], font=Font(strike=True, color="808080")))
    dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
    dv.add(f"M{FIRST}:M{end}")
    ws.add_data_validation(dv)
    dvw = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                         error="Enter wastage as a percentage, e.g. 5%")
    dvw.add(f"I{FIRST}:I{end}")
    ws.add_data_validation(dvw)

    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:Q{end}"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&8Page &P of &N"
    ws.oddFooter.left.text = f"&8{c['title']}"
    c["end"] = end

# ---------------------------------------------------------------- ALL ITEMS (auto-consolidated)
db = wb.create_sheet("ALL ITEMS", 1)
db.sheet_properties.tabColor = NAVY
db.sheet_view.showGridLines = False
DBH = ["Category", "Item Code", "Sub-Category", "Item Description", "Specification / Size", "Finish / Colour",
       "Brand / Origin", "Unit", "Net Cost (AED)", "Selling (AED)", "Selling incl. VAT", "Rate Status", "Search key"]
db.merge_cells("A1:L1")
db["A1"] = "ALL ITEMS - MASTER LIST (auto-linked from category sheets; do not edit here)"
db["A1"].font = F(bold=True, size=14, color="FFFFFF")
db["A1"].fill = HDR_FILL
db["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
db.row_dimensions[1].height = 28
for j, h in enumerate(DBH, start=1):
    cell = db.cell(2, j, h)
    cell.font = F(bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER
SRC = {2: "A", 3: "B", 4: "C", 5: "D", 6: "E", 7: "F", 8: "G", 9: "J", 10: "K", 11: "L", 12: "M"}
for i, (sheet, title, r) in enumerate(DB_ROWS, start=3):
    q = f"'{sheet}'!"
    db.cell(i, 1, f'=IF({q}$C${r}="","","{title}")')
    for j, col in SRC.items():
        db.cell(i, j, f'=IF({q}$C${r}="","",{q}{col}{r})')
    db.cell(i, 13, (f'=IF(D{i}="","",IF(AND(OR(Search!$C$6="All Categories",A{i}=Search!$C$6),'
                    f'OR(Search!$P$5="",ISNUMBER(SEARCH(" "&Search!$P$5,N{i}))),OR(Search!$P$6="",ISNUMBER(SEARCH(" "&Search!$P$6,N{i}))),OR(Search!$P$7="",ISNUMBER(SEARCH(" "&Search!$P$7,N{i}))),OR(Search!$P$8="",ISNUMBER(SEARCH(" "&Search!$P$8,N{i}))),OR(Search!$P$9="",ISNUMBER(SEARCH(" "&Search!$P$9,N{i}))),OR(Search!$P$10="",ISNUMBER(SEARCH(" "&Search!$P$10,N{i})))),ROW(),""))'))
    raw = f'B{i}&" "&C{i}&" "&D{i}&" "&E{i}&" "&F{i}&" "&G{i}&" "&H{i}'
    for ch in ("(", ")", "/", "-", ","):
        raw = f'SUBSTITUTE({raw},"{ch}"," ")'
    db.cell(i, 14, f'=IF(D{i}="",""," "&{raw}&" ")')
    for j in range(1, 14):
        cell = db.cell(i, j)
        cell.font = F(size=9, color="008000" if j < 13 else "808080")
        cell.alignment = LEFT if j in (1, 3, 4, 5, 6, 7) else CENTER
        if j in (9, 10, 11):
            cell.number_format = NUM
DB_END = 2 + len(DB_ROWS)
db.conditional_formatting.add(f"A3:L{DB_END}", FormulaRule(
    formula=['AND($D3<>"",MOD(ROW(),2)=0)'], fill=PatternFill("solid", fgColor=GREY)))
for j, w in enumerate([30, 10, 22, 40, 30, 18, 24, 7, 12, 12, 13, 11, 9], start=1):
    db.column_dimensions[get_column_letter(j)].width = w
db.column_dimensions["M"].hidden = True
db.column_dimensions["N"].hidden = True
db.freeze_panes = "E3"
db.auto_filter.ref = f"A2:L{DB_END}"

# ---------------------------------------------------------------- Search
sr = wb.create_sheet("Search", 1)
sr.sheet_properties.tabColor = "FFC000"
sr.sheet_view.showGridLines = False
sr.merge_cells("B2:L2")
sr["B2"] = "QUICK MATERIAL SEARCH"
sr["B2"].font = F(bold=True, size=18, color="FFFFFF")
sr["B2"].fill = HDR_FILL
sr["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
sr.row_dimensions[2].height = 36
sr.merge_cells("B3:L3")
sr["B3"] = ("Type one or more words in any order (e.g. aluminium sheet, 12 tempered, floor spring, 316 tube, "
            "DGU low e) - every word must match the start of a word. Pick a category if needed. Results update instantly.")
sr["B3"].font = F(italic=True, size=9, color="404040")
for r, lbl in ((5, "Search keyword:"), (6, "Category:"), (7, "Items found:")):
    sr.merge_cells(f"A{r}:B{r}")
    sr[f"A{r}"] = lbl
    sr[f"A{r}"].font = F(bold=True, size=11, color=NAVY)
    sr[f"A{r}"].alignment = Alignment(horizontal="right", vertical="center")
sr.merge_cells("C5:E5")
sr.merge_cells("C6:E6")
sr["C5"] = "aluminium sheet"
sr["C6"] = "All Categories"
for a in ("C5", "C6"):
    sr[a].font = F(bold=True, size=12, color="0000FF")
    sr[a].fill = PatternFill("solid", fgColor="FFFF00")
    sr[a].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for col in "CDE":
    for r in (5, 6):
        sr[f"{col}{r}"].border = Border(left=Side(style="medium", color=NAVY), right=Side(style="medium", color=NAVY),
                                        top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
sr.row_dimensions[5].height = 24
sr.row_dimensions[6].height = 24
sr["C7"] = f"=COUNT('ALL ITEMS'!M3:M{DB_END})"
sr["C7"].font = F(bold=True, size=12, color="C00000")
sr["C7"].alignment = Alignment(horizontal="left", indent=1)
sr["G5"] = "Tip: clear the keyword to list a whole category."
sr["G6"] = "Max 300 results shown - refine keyword if needed."
for a in ("G5", "G6"):
    sr[a].font = F(italic=True, size=9, color="595959")
sr["G7"] = '=HYPERLINK("#Index!A1","<< Back to Index")'
sr["G7"].font = F(size=9, color="0563C1", underline="single")

SH = ["#", "Item Code", "Category", "Item Description", "Specification / Size", "Finish / Colour",
      "Brand / Origin", "Unit", "Net Cost (AED)", "Selling (AED)", "Selling incl. VAT", "Rate Status"]
SMAP = {"B": 2, "C": 1, "D": 4, "E": 5, "F": 6, "G": 7, "H": 8, "I": 9, "J": 10, "K": 11, "L": 12}
HR = 9
for j, h in enumerate(SH, start=1):
    cell = sr.cell(HR, j, h)
    cell.font = F(bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER
N_RES = 300
for k in range(N_RES):
    r = HR + 1 + k
    sr.cell(r, 14, f"=IF(ROWS($N${HR + 1}:N{r})>$C$7,\"\",SMALL('ALL ITEMS'!$M$3:$M${DB_END},ROWS($N${HR + 1}:N{r})))")
    sr.cell(r, 1, f'=IF($N{r}="","",ROWS($N${HR + 1}:N{r}))')
    for col, src in SMAP.items():
        sr[f"{col}{r}"] = f"=IF($N{r}=\"\",\"\",INDEX('ALL ITEMS'!${get_column_letter(src)}:${get_column_letter(src)},$N{r}))"
    for j in range(1, 13):
        cell = sr.cell(r, j)
        cell.font = F(size=9)
        cell.alignment = LEFT if j in (3, 4, 5, 6, 7) else CENTER
        if j in (9, 10, 11):
            cell.number_format = NUM
sr.conditional_formatting.add(f"A{HR + 1}:L{HR + N_RES}", FormulaRule(
    formula=[f'$N{HR + 1}<>""'], border=BORDER))
sr.conditional_formatting.add(f"A{HR + 1}:L{HR + N_RES}", FormulaRule(
    formula=[f'AND($N{HR + 1}<>"",MOD(ROW(),2)=0)'], fill=PatternFill("solid", fgColor=GREY)))
sr.conditional_formatting.add(f"I{HR + 1}:K{HR + N_RES}", FormulaRule(
    formula=[f'$N{HR + 1}<>""'], font=Font(bold=True, color=NAVY)))
for j, w in enumerate([5, 10, 28, 40, 30, 18, 24, 7, 12, 12, 13, 11, 2, 6], start=1):
    sr.column_dimensions[get_column_letter(j)].width = w
sr.column_dimensions["N"].hidden = True
# keyword split into words (P = word, Q = remaining text): every word must match, any order
sr["Q4"] = '=TRIM($C$5)&" "'
for r in range(5, 11):
    sr[f"P{r}"] = f'=IFERROR(LEFT(Q{r - 1},FIND(" ",Q{r - 1})-1),"")'
    sr[f"Q{r}"] = f'=IFERROR(MID(Q{r - 1},FIND(" ",Q{r - 1})+1,999),"")'
sr.column_dimensions["P"].hidden = True
sr.column_dimensions["Q"].hidden = True
sr.merge_cells("A8:B8")
sr["A8"] = "Words matched:"
sr["A8"].font = F(size=9, italic=True, color="595959")
sr["A8"].alignment = Alignment(horizontal="right")
sr["C8"] = '=IF(P5="","(none - showing all)",P5&IF(P6="",""," + "&P6)&IF(P7="",""," + "&P7)&IF(P8="",""," + "&P8)&IF(P9="",""," + "&P9)&IF(P10="",""," + "&P10))'
sr["C8"].font = F(size=9, italic=True, color="595959")
sr.freeze_panes = f"A{HR + 1}"

# ---------------------------------------------------------------- Suppliers
sp = wb.create_sheet("Suppliers")
sp.sheet_properties.tabColor = "70AD47"
sp.sheet_view.showGridLines = False
sp.merge_cells("A1:K1")
sp["A1"] = "SUPPLIER DIRECTORY"
sp["A1"].font = F(bold=True, size=14, color="FFFFFF")
sp["A1"].fill = HDR_FILL
sp.row_dimensions[1].height = 28
SPH = ["Sr.", "Category", "Supplier", "Products / Scope", "Contact Person", "Mobile / Tel.", "Email",
       "Payment Terms", "Lead Time", "Rating (1-5)", "Remarks"]
for j, h in enumerate(SPH, start=1):
    cell = sp.cell(3, j, h)
    cell.font = F(bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER
for i in range(len(SUPPLIERS) + 20):
    r = 4 + i
    if i < len(SUPPLIERS):
        catg, name, prod = SUPPLIERS[i]
        sp.cell(r, 2, catg)
        sp.cell(r, 3, name)
        sp.cell(r, 4, prod)
    sp.cell(r, 1, f'=IF(C{r}="","",ROW()-3)')
    for j in range(1, 12):
        cell = sp.cell(r, j)
        cell.border = BORDER
        cell.font = F(size=10, color="0000FF" if j > 1 else "000000")
        cell.alignment = LEFT if j > 1 else CENTER
        if i >= len(SUPPLIERS) or j >= 5:
            cell.fill = INPUT_FILL if j > 1 else PatternFill()
for j, w in enumerate([5, 22, 40, 40, 20, 16, 26, 16, 12, 11, 24], start=1):
    sp.column_dimensions[get_column_letter(j)].width = w
sp.freeze_panes = "D4"
sp.auto_filter.ref = f"A3:K{3 + len(SUPPLIERS) + 20}"
dvr = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
dvr.add(f"J4:J{3 + len(SUPPLIERS) + 20}")
sp.add_data_validation(dvr)

# ---------------------------------------------------------------- Rate Revision Log
lg = wb.create_sheet("Rate Log")
lg.sheet_properties.tabColor = "7F7F7F"
lg.sheet_view.showGridLines = False
lg.merge_cells("A1:K1")
lg["A1"] = "RATE REVISION LOG"
lg["A1"].font = F(bold=True, size=14, color="FFFFFF")
lg["A1"].fill = HDR_FILL
lg.row_dimensions[1].height = 28
LGH = ["Date", "Item Code", "Item Description (auto)", "Old Rate (AED)", "New Rate (AED)", "Change %",
       "Supplier", "Quote Ref.", "Reason", "Updated By", "Approved By"]
for j, h in enumerate(LGH, start=1):
    cell = lg.cell(3, j, h)
    cell.font = F(bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER
for r in range(4, 104):
    lg.cell(r, 3, f"=IF(B{r}=\"\",\"\",IFERROR(INDEX('ALL ITEMS'!$D$3:$D${DB_END},MATCH(B{r},'ALL ITEMS'!$B$3:$B${DB_END},0)),\"Code not found\"))")
    lg.cell(r, 6, f'=IF(OR(D{r}="",E{r}="",D{r}=0),"",(E{r}-D{r})/D{r})')
    for j in range(1, 12):
        cell = lg.cell(r, j)
        cell.border = BORDER
        cell.font = F(size=10, color="000000" if j in (3, 6) else "0000FF")
        cell.alignment = LEFT if j in (3, 7, 9) else CENTER
        if j in (4, 5):
            cell.number_format = NUM
        if j == 6:
            cell.number_format = '+0.0%;-0.0%;"-"'
        if j == 1:
            cell.number_format = "dd-mmm-yy"
# example row
lg["A4"] = "09-Oct-2026"
lg["B4"], lg["D4"], lg["E4"] = "GL-010", 115.0, 120.0
lg["G4"], lg["H4"], lg["I4"], lg["J4"] = "Example supplier", "Q-0001", "Example entry - replace", "Estimator"
lg.conditional_formatting.add("F4:F103", FormulaRule(formula=['AND(F4<>"",F4>0)'], font=Font(color="C00000", bold=True)))
lg.conditional_formatting.add("F4:F103", FormulaRule(formula=['AND(F4<>"",F4<0)'], font=Font(color="008000", bold=True)))
for j, w in enumerate([11, 10, 40, 13, 13, 10, 22, 13, 30, 14, 14], start=1):
    lg.column_dimensions[get_column_letter(j)].width = w
lg.freeze_panes = "A4"

# ---------------------------------------------------------------- Index (dashboard)
ix = ws0
ix.sheet_properties.tabColor = NAVY
ix.sheet_view.showGridLines = False
ix.merge_cells("B2:H2")
ix["B2"] = "MATERIAL PRICE LIST & RATE DATABASE"
ix["B2"].font = F(bold=True, size=20, color="FFFFFF")
ix["B2"].fill = HDR_FILL
ix["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ix.row_dimensions[2].height = 42
ix.merge_cells("B3:H3")
ix["B3"] = "Saeed Al Siraj Glass & Aluminium Works L.L.C.  |  Estimation Department  |  Aluminium, Glass, MS & SS Works"
ix["B3"].font = F(bold=True, size=10, color=NAVY)
ix["B3"].fill = PatternFill("solid", fgColor=GOLD)
ix["B3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ix.row_dimensions[3].height = 22

info = [("Rate validity / last reviewed:", "=Settings!C%d" % (last + 1)),
        ("VAT:", "=VAT"), ("Default O&P:", "=OHP"),
        ("Total items:", f"=COUNTA('ALL ITEMS'!D3:D{DB_END})-COUNTBLANK('ALL ITEMS'!D3:D{DB_END})")]
for i, (k, v) in enumerate(info, start=5):
    ix[f"B{i}"] = k
    ix[f"B{i}"].font = F(bold=True, size=10, color="404040")
    ix[f"D{i}"] = v
    ix[f"D{i}"].font = F(bold=True, size=10, color=NAVY)
    ix[f"D{i}"].alignment = Alignment(horizontal="left")
ix["D6"].number_format = ix["D7"].number_format = "0%"
ix["F5"] = '=HYPERLINK("#Search!C5",">>  SEARCH MATERIAL  <<")'
ix["F5"].font = F(bold=True, size=12, color="FFFFFF")
ix["F5"].fill = PatternFill("solid", fgColor="C00000")
ix["F5"].alignment = CENTER
ix.merge_cells("F5:G6")
ix["F7"] = '=HYPERLINK("#Settings!C6","Edit base rates / VAT / O&P")'
ix["F7"].font = F(size=9, color="0563C1", underline="single")
ix["F8"] = '=HYPERLINK("#\'ALL ITEMS\'!A1","View master list (all items)")'
ix["F8"].font = F(size=9, color="0563C1", underline="single")

R0 = 10
for j, h in enumerate(["No.", "Category", "Code", "Scope / Contents", "Items", "Open", "Sheet"], start=2):
    cell = ix.cell(R0, j, h)
    cell.font = F(bold=True, color="FFFFFF", size=10)
    cell.fill = HDR_FILL
    cell.alignment = CENTER
    cell.border = BORDER
for i, c in enumerate(CATS, start=1):
    r = R0 + i
    ix.cell(r, 2, i)
    ix.cell(r, 3, c["title"])
    ix.cell(r, 4, c["prefix"])
    ix.cell(r, 5, c["scope"])
    ix.cell(r, 6, f"=COUNTA('{c['sheet']}'!C5:C{c['end']})")
    ix.cell(r, 7, f'=HYPERLINK("#\'{c["sheet"]}\'!A1","Open  >")')
    ix.cell(r, 8, c["sheet"])
    for j in range(2, 9):
        cell = ix.cell(r, j)
        cell.border = BORDER
        cell.font = F(size=10, bold=(j == 3))
        cell.alignment = LEFT if j in (3, 5) else CENTER
    ix.cell(r, 2).fill = PatternFill("solid", fgColor=c["colour"])
    ix.cell(r, 7).font = F(size=10, bold=True, color="0563C1", underline="single")
    ix.cell(r, 8).font = F(size=8, color="808080")
rt = R0 + len(CATS) + 1
ix.cell(rt, 3, "TOTAL")
ix.cell(rt, 6, f"=SUM(F{R0 + 1}:F{rt - 1})")
for j in range(2, 9):
    cell = ix.cell(rt, j)
    cell.font = F(bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.alignment = CENTER if j != 3 else LEFT
    cell.border = BORDER

# category list for Search dropdown (column K, hidden)
ix["K10"] = "All Categories"
for i, c in enumerate(CATS, start=1):
    ix[f"K{10 + i}"] = c["title"]
ix.column_dimensions["K"].hidden = True
dvc = DataValidation(type="list", formula1=f"=Index!$K$10:$K${10 + len(CATS)}", allow_blank=False)
dvc.add("C6")
sr.add_data_validation(dvc)

# How to use + legend
hu = rt + 2
ix[f"B{hu}"] = "HOW TO USE"
ix[f"B{hu}"].font = F(bold=True, size=12, color=NAVY)
steps = [
    "1.  FIND: Go to 'Search' - type any keyword (size, brand, item) - results with rates appear instantly.",
    "2.  BROWSE: Click 'Open' above to go to a category sheet. Use the filter arrows on the header row.",
    "3.  UPDATE RATE: Type the new Basic Rate (blue) in the category sheet, set Rate Status, Supplier, Quote Ref & Date, then log it in 'Rate Log'.",
    "4.  ADD ITEM: Use the yellow spare rows at the bottom of each category sheet - the code, net, selling & VAT rates fill in automatically.",
    "5.  MARKET CHANGE: Change aluminium / steel / SS rate per kg, VAT or O&P in 'Settings' - all linked rates (green) update automatically.",
    "6.  NET COST RATE = Basic Rate x (1 + Wastage %).  SELLING = Net Cost x (1 + O&P).  SELLING incl. VAT = Selling x (1 + VAT).",
]
for i, s in enumerate(steps, start=hu + 1):
    ix.merge_cells(f"B{i}:H{i}")
    ix[f"B{i}"] = s
    ix[f"B{i}"].font = F(size=10)
    ix[f"B{i}"].alignment = LEFT
    ix.row_dimensions[i].height = 18

lgd = hu + len(steps) + 2
ix[f"B{lgd}"] = "LEGEND"
ix[f"B{lgd}"].font = F(bold=True, size=12, color=NAVY)
legend = [
    ("Blue text", "Input / editable value", Font(name=FONT, color="0000FF", bold=True), None),
    ("Black text", "Formula - do not overwrite", Font(name=FONT, color="000000", bold=True), None),
    ("Green text", "Linked to Settings / another sheet", Font(name=FONT, color="008000", bold=True), None),
    ("Indicative", "Budget market rate - confirm before final quote", None, "FCE4D6"),
    ("Quoted", "Rate taken from a supplier quotation", None, "DDEBF7"),
    ("Verified", "Rate confirmed by recent purchase / LPO", None, "E2EFDA"),
    ("Obsolete", "Item no longer used (shown struck-through)", None, "D9D9D9"),
]
for i, (k, v, fnt, fill) in enumerate(legend, start=lgd + 1):
    ix[f"C{i}"] = k
    ix[f"C{i}"].font = fnt or F(size=10, bold=True)
    if fill:
        ix[f"C{i}"].fill = PatternFill("solid", fgColor=fill)
    ix[f"C{i}"].border = BORDER
    ix[f"C{i}"].alignment = CENTER
    ix[f"D{i}"] = v
    ix[f"D{i}"].font = F(size=10)
    ix.merge_cells(f"D{i}:F{i}")

un = lgd + len(legend) + 2
ix[f"B{un}"] = "UNITS"
ix[f"B{un}"].font = F(bold=True, size=12, color=NAVY)
units = "m2 = square metre  |  lm = linear metre  |  kg = kilogram  |  nos = numbers  |  set / pair / pkt / box / roll  |  ltr = litre  |  day / hr / month / trip  |  ls = lump sum  |  % = percentage"
ix.merge_cells(f"B{un + 1}:H{un + 1}")
ix[f"B{un + 1}"] = units
ix[f"B{un + 1}"].font = F(size=9)
ix[f"B{un + 1}"].alignment = LEFT

dis = un + 3
ix.merge_cells(f"B{dis}:H{dis + 1}")
ix[f"B{dis}"] = ("DISCLAIMER: Rates are indicative UAE market budget rates (Oct-2026) for estimation reference. "
                 "Always confirm with current supplier quotations, project specification and quantities before "
                 "submitting a final quotation.")
ix[f"B{dis}"].font = F(italic=True, size=9, color="C00000")
ix[f"B{dis}"].alignment = LEFT

for col, w in zip("ABCDEFGH", [2, 6, 44, 8, 70, 9, 12, 20]):
    ix.column_dimensions[col].width = w

for w in (ix, sr, db, sp, lg, st):
    w.page_setup.orientation = "landscape"
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
sr.print_area = f"A1:L{HR + N_RES}"
ix.print_area = f"A1:H{dis + 1}"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "items:", sum(len(c["items"]) for c in CATS))
